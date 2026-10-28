"""Полный HTTP-путь на 127.0.0.1: загрузка файлов, импорт, скачивание."""
import hashlib
import http.client
import io
import json
import re
import sys
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'vendor'), str(ROOT)]
from openpyxl import load_workbook
from app.web import AppState, LocalServer, MAX_FILE
from app.importer import import_files


class WebTests(unittest.TestCase):
    def setUp(self):
        self.state = AppState()
        self.server = LocalServer(0, self.state)
        self.thread = threading.Thread(target=self.server.serve_forever, kwargs={'poll_interval': 0.01})
        self.thread.start()
        self.matrix = ROOT / 'examples/01_matrix.xlsx'
        self.templates = [ROOT / 'examples/02_template_update.xlsx', ROOT / 'examples/03_template_new_group.xlsx']

    def tearDown(self):
        self.server.shutdown()
        self.thread.join()
        self.server.server_close()
        folder = Path(self.state.temp.name)
        self.state.close()
        self.assertFalse(folder.exists())

    def request(self, method, path, body=None, headers=None, token=True):
        headers = dict(headers or {})
        if token:
            headers['X-Matrix-Token'] = self.state.token
        if isinstance(body, dict):
            body = json.dumps(body).encode()
            headers['Content-Type'] = 'application/json'
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=10)
        connection.request(method, path, body=body, headers=headers)
        response = connection.getresponse()
        data = response.read()
        result = (response.status, data, dict(response.getheaders()))
        connection.close()
        return result

    def json_request(self, method, path, body=None, **kwargs):
        code, data, _ = self.request(method, path, body, **kwargs)
        return code, json.loads(data)

    def create(self, matrix=None, templates=None, environment='IFT'):
        matrix = matrix or self.matrix.read_bytes()
        templates = templates or [p.read_bytes() for p in self.templates]
        files = [matrix, *templates]
        manifest = [{'name': f'потреба_{i}.xlsx', 'size':len(data)} for i, data in enumerate(files)]
        code, job = self.json_request('POST', '/api/jobs', {'matrix':manifest[0], 'templates':manifest[1:], 'environment':environment})
        self.assertEqual(code, 201, job)
        ident = job['id']
        for index, data in enumerate(files):
            code, response = self.json_request('PUT', f'/api/jobs/{ident}/files/{index}', data)
            self.assertEqual(code, 200, response)
        return ident

    def wait(self, ident):
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            code, job = self.json_request('GET', f'/api/jobs/{ident}')
            self.assertEqual(code, 200)
            if job['state'] in ('done','error'):
                return job
            time.sleep(0.01)
        self.fail('Импорт не завершился')

    def test_page_assets_and_no_directory_listing(self):
        code, html, headers = self.request('GET', '/', token=False)
        self.assertEqual(code,200)
        self.assertIn(self.state.token.encode(),html)
        self.assertIn('frame-ancestors',headers['Content-Security-Policy'])
        for asset in ['/app.js','/style.css']:
            self.assertEqual(self.request('GET',asset)[0],200)
        for path in ['/config.json','/vendor/','/../../README.md']:
            self.assertEqual(self.request('GET',path)[0],404)
        self.assertEqual(self.server.server_address[0],'127.0.0.1')

    def test_upload_import_download_and_repeat(self):
        original = hashlib.sha256(self.matrix.read_bytes()).hexdigest()
        ident = self.create()
        self.assertEqual(self.json_request('POST',f'/api/jobs/{ident}/start',{})[0],202)
        job = self.wait(ident)
        self.assertEqual(job['state'],'done',job)
        self.assertEqual(job['counts'],{'added':4,'updated':0,'skipped':1,'conflicted':0,'versioned':1,'promoted':0,'prom_deactivated':0,'total':5})
        self.assertEqual(job['environment'], 'IFT')
        self.assertEqual(job['conflicts'], [])
        code, payload, headers = self.request('GET',job['download_url'])
        self.assertEqual(code,200)
        self.assertIn('attachment',headers['Content-Disposition'])
        workbook = load_workbook(io.BytesIO(payload))
        self.assertEqual(workbook.worksheets[0].max_row,9)
        self.assertIsNone(workbook.worksheets[0]['AL6'].value)
        self.assertIsNone(workbook.worksheets[0]['C2'].value)
        self.assertEqual(workbook.worksheets[0]['D2'].value, 1)
        self.assertEqual(workbook.worksheets[0]['C6'].value, 1)
        self.assertIsNone(workbook.worksheets[0]['D6'].value)
        self.assertEqual(workbook.worksheets[0]['A6'].value,'=ROW(A7)-2')
        workbook.close()
        second = self.create(matrix=payload)
        self.json_request('POST',f'/api/jobs/{second}/start',{})
        self.assertEqual(self.wait(second)['counts'],{'added':0,'updated':0,'skipped':5,'conflicted':0,'versioned':0,'promoted':0,'prom_deactivated':0,'total':5})
        self.assertEqual(hashlib.sha256(self.matrix.read_bytes()).hexdigest(),original)

    def test_demo(self):
        self.assertTrue(self.json_request('GET','/api/info')[1]['demo'])
        code, response = self.json_request('POST','/api/demo',{})
        self.assertEqual(code,202)
        job = self.wait(response['id'])
        self.assertEqual(job['counts']['added'],4)

    def test_partial_result_has_conflict_details_and_can_be_downloaded(self):
        book = load_workbook(self.matrix)
        sheet = book.worksheets[0]
        sheet.append([c.value for c in sheet[2]])
        original = [c.value for c in sheet[2]]
        data = io.BytesIO()
        book.save(data)
        book.close()
        ident = self.create(matrix=data.getvalue())
        self.json_request('POST', f'/api/jobs/{ident}/start', {})
        job = self.wait(ident)
        self.assertEqual(job['state'], 'done', job)
        self.assertEqual(job['counts'], {'added':3, 'updated':0, 'skipped':1, 'conflicted':1, 'versioned':0, 'promoted':0,'prom_deactivated':0,'total':5})
        self.assertEqual(len(job['conflicts']), 1)
        self.assertIn('[2, 6]', job['conflicts'][0])
        code, payload, _ = self.request('GET', job['download_url'])
        self.assertEqual(code, 200)
        result = load_workbook(io.BytesIO(payload))
        self.assertEqual([c.value for c in result.worksheets[0][2]], original)
        self.assertEqual(result.worksheets[0].max_row, 9)
        result.close()

    def test_api_requires_token(self):
        self.assertEqual(self.request('GET','/api/info',token=False)[0],403)
        self.assertEqual(self.request('POST','/api/demo',{},token=False)[0],403)

    def test_foreign_host_and_origin_are_rejected(self):
        self.assertEqual(self.request('GET','/',headers={'Host':'attacker.invalid'})[0],403)
        self.assertEqual(self.request('POST','/api/demo',{},headers={'Origin':'https://attacker.invalid'})[0],403)

    def test_incomplete_upload_cannot_start(self):
        data={'name':'matrix.xlsx','size':10}
        _, job = self.json_request('POST','/api/jobs',{'matrix':data,'templates':[data], 'environment':'IFT'})
        code, _ = self.json_request('POST',f'/api/jobs/{job["id"]}/start',{})
        self.assertEqual(code,409)
        self.assertEqual(self.request('GET',f'/api/jobs/{job["id"]}/download')[0],409)

    def test_invalid_names_and_sizes(self):
        valid={'name':'ok.xlsx','size':1}
        for bad in [{'name':'../bad.xlsx','size':1}, {'name':'file.xls','size':1}, {'name':'file.xlsx','size':0}, {'name':'file.xlsx','size':MAX_FILE+1}]:
            code,_=self.json_request('POST','/api/jobs',{'matrix':bad,'templates':[valid], 'environment':'IFT'})
            self.assertEqual(code,400)

    def test_wrong_upload_size_and_duplicate_upload(self):
        item={'name':'ok.xlsx','size':4}
        _,job=self.json_request('POST','/api/jobs',{'matrix':item,'templates':[item], 'environment':'IFT'})
        path=f'/api/jobs/{job["id"]}/files/0'
        self.assertEqual(self.request('PUT',path,b'123')[0],400)
        self.assertEqual(self.request('PUT',path,b'1234')[0],200)
        self.assertEqual(self.request('PUT',path,b'1234')[0],409)

    def test_corrupt_xlsx_returns_error_without_download(self):
        ident = self.create(matrix=b'not a zip')
        self.json_request('POST',f'/api/jobs/{ident}/start',{})
        job = self.wait(ident)
        self.assertEqual(job['state'],'error')
        self.assertIn('повреждён',job['error'])
        self.assertEqual(self.request('GET',f'/api/jobs/{ident}/download')[0],409)

    def test_default_launch_does_not_import_tkinter(self):
        import run
        with patch.dict(sys.modules,{'tkinter':None,'_tkinter':None}), patch.object(sys,'argv',['run.py','--no-browser']), patch('app.web.start') as start:
            self.assertEqual(run.main(),0)
            start.assert_called_once_with(None,8765,open_browser=False)

    def test_environment_is_required_for_real_uploads(self):
        item = {'name':'ok.xlsx', 'size':1}
        for environment in (None, '', 'BOTH', 1):
            code, response = self.json_request('POST', '/api/jobs', {'matrix':item, 'templates':[item], 'environment':environment})
            self.assertEqual(code, 400)
            self.assertIn('среду', response['error'])

    def test_prom_import_activates_ready_ift_without_adding_rows(self):
        prepared = Path(self.state.temp.name) / 'ready_ift.xlsx'
        import_files(self.matrix, self.templates, prepared, environment='IFT')
        ident = self.create(matrix=prepared.read_bytes(), environment='PROM')
        self.json_request('POST', f'/api/jobs/{ident}/start', {})
        job = self.wait(ident)
        self.assertEqual(job['state'], 'done', job)
        self.assertEqual(job['environment'], 'PROM')
        self.assertEqual(job['counts'], {'added':0, 'updated':5, 'skipped':0, 'conflicted':0, 'versioned':0, 'promoted':5, 'prom_deactivated':1,'total':5})
        code, payload, headers = self.request('GET', job['download_url'])
        self.assertEqual(code, 200)
        self.assertIn('_PROM_', headers['Content-Disposition'])
        book = load_workbook(io.BytesIO(payload))
        sheet = book.worksheets[0]
        self.assertIsNone(sheet['C2'].value)
        self.assertIsNone(sheet['D2'].value)
        self.assertEqual(sheet['C6'].value, 1)
        self.assertEqual(sheet['D6'].value, 1)
        self.assertEqual(sheet.max_row, 9)
        book.close()
        second = self.create(matrix=payload, environment='PROM')
        self.json_request('POST', f'/api/jobs/{second}/start', {})
        repeat = self.wait(second)
        self.assertEqual((repeat['counts']['added'], repeat['counts']['updated'], repeat['counts']['skipped']), (0, 0, 5))

    def test_cli_requires_explicit_environment(self):
        import run
        with patch.object(sys, 'argv', ['run.py', '--matrix', 'm.xlsx', '--templates', 't.xlsx', '--output', 'o.xlsx']):
            with patch('sys.stderr', new=io.StringIO()), self.assertRaises(SystemExit) as error:
                run.main()
            self.assertEqual(error.exception.code, 2)


if __name__ == '__main__':
    unittest.main(verbosity=2)
