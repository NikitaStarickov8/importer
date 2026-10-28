"""Среды, версии, неизменность ключей, ручной байпас и тип транзакции."""
import json
import unittest

import test_importer as importer_tests
from app.importer import Catalog, ImportProblem, columns_at, load_config, read_matrix_row
from openpyxl import load_workbook


class WorkflowTests(unittest.TestCase):
    setUp = importer_tests.ImportTests.setUp
    tearDown = importer_tests.ImportTests.tearDown
    change = importer_tests.ImportTests.change
    run_import = importer_tests.ImportTests.run_import

    def single_template(self, edit=lambda sheet: None):
        def modify(book):
            sheet = book.worksheets[0]
            for row in (19, 20):
                for cell in sheet[row]:
                    cell.value = None
            edit(sheet)
        return self.change(self.templates[0], modify, 'single.xlsx')

    def matrix_with_prom(self, ift=1):
        def edit(book):
            sheet = book.worksheets[0]
            sheet['C2'], sheet['D2'], sheet['AL2'] = ift, 1, 'ТЗ'
            sheet['AN2'] = 'Служебное поле не менять'
        return self.change(self.matrix, edit, 'prod.xlsx')

    def check_version(self, environment, ift=1):
        matrix = self.matrix_with_prom(ift)
        result = self.run_import(matrix=matrix, templates=[self.single_template()], environment=environment)
        self.assertEqual((result.added, result.updated, result.skipped, result.versioned), (1, 0, 0, 1))
        self.assertEqual(result.environment, environment)
        before, after = load_workbook(matrix), load_workbook(self.output)
        try:
            a, b = before.worksheets[0], after.worksheets[0]
            expected = [c.value for c in a[2]]
            expected[2 if environment == 'IFT' else 3] = None
            self.assertEqual([c.value for c in b[2]], expected)
            self.assertEqual(b['C6'].value, 1 if environment == 'IFT' else None)
            self.assertEqual(b['D6'].value, 1 if environment == 'PROM' else None)
            self.assertEqual(b['AK6'].value, 'Новое тестовое сообщение. Повторите операцию позже.')
            self.assertEqual(b['AN6'].value, 'Служебное поле не менять')
            self.assertEqual(b['B6'].value, a['B2'].value)
            self.assertEqual(b['AE6'].value, a['AE2'].value)
            self.assertEqual(b['AJ6'].value, a['AJ2'].value)
            self.assertEqual(b['AL2'].value, 'ТЗ')
            self.assertIsNone(b['AL6'].value)
            self.assertEqual(b['M6'].value, 'Не финансовые операции')
            self.assertEqual(b['A6'].value, '=ROW(A7)-2')
            self.assertIn('MATCH(G6,', b['H6'].value)
            self.assertEqual(a['AK2']._style, b['AK6']._style)
            self.assertEqual(b.tables['MatrixTable'].ref, 'A1:AQ6')
        finally:
            before.close()
            after.close()

    def test_ift_splits_shared_row_and_preserves_prom(self):
        self.check_version('IFT')

    def test_prom_does_not_rewrite_unprepared_shared_row(self):
        result = self.run_import(matrix=self.matrix_with_prom(), templates=[self.single_template()], environment='PROM')
        self.assertEqual((result.added, result.updated), (0, 0))
        self.assertEqual(len(result.conflicts), 1)

    def test_prom_without_ready_ift_is_manual(self):
        result = self.run_import(matrix=self.matrix_with_prom(ift=None), templates=[self.single_template()], environment='PROM')
        self.assertEqual((result.added, result.updated), (0, 0))
        self.assertEqual(len(result.conflicts), 1)

    def test_ift_only_update_changes_only_three_payload_fields(self):
        def edit(book):
            sheet = book.worksheets[0]
            sheet['G2'], sheet['AL2'] = ' demo_loan ', 'ТЗ'
        matrix = self.change(self.matrix, edit, 'keys.xlsx')
        def change_payload(sheet):
            sheet['L18'], sheet['O18'], sheet['P18'], sheet['Q18'] = 'З', 'Новый заголовок', 'Новый текст', 'произвольный байпас'
        result = self.run_import(matrix=matrix, templates=[self.single_template(change_payload)])
        self.assertEqual((result.added, result.updated, result.versioned), (0, 1, 0))
        before, after = load_workbook(matrix), load_workbook(self.output)
        try:
            a, b = before.worksheets[0], after.worksheets[0]
            expected = [c.value for c in a[2]]
            expected[30], expected[35], expected[36] = 'З', 'Новый заголовок', 'Новый текст'
            self.assertEqual([c.value for c in b[2]], expected)
            self.assertEqual(b.max_row, a.max_row)
        finally:
            before.close()
            after.close()

    def test_identical_prod_payload_does_not_create_version(self):
        matrix = self.matrix_with_prom()
        source = load_workbook(matrix)
        text = source.worksheets[0]['AK2'].value
        source.close()
        template = self.single_template(lambda s: setattr(s['P18'], 'value', text))
        for environment in ('IFT', 'PROM'):
            output = self.folder / f'identical_{environment}.xlsx'
            result = self.run_import(matrix=matrix, templates=[template], environment=environment, output=output)
            self.assertEqual((result.added, result.updated, result.skipped, result.versioned), (0, 0, 1, 0))
            book = load_workbook(output)
            self.assertEqual((book.worksheets[0]['C2'].value, book.worksheets[0]['D2'].value), (1, 1))
            book.close()

    def test_repeated_versions_are_idempotent_in_both_environments(self):
        template, matrix = self.single_template(), self.matrix_with_prom()
        for environment in ('IFT', 'PROM'):
            first = self.folder / f'first_{environment}.xlsx'
            source = matrix
            if environment == 'PROM':
                source = self.folder / 'prepared_ift.xlsx'
                self.run_import(matrix=matrix, templates=[template], environment='IFT', output=source)
            self.run_import(matrix=source, templates=[template], environment=environment, output=first)
            result = self.run_import(matrix=first, templates=[template], environment=environment, output=self.folder / f'second_{environment}.xlsx')
            self.assertEqual((result.added, result.updated, result.skipped, result.versioned), (0, 0, 1, 0))

    def test_ift_then_prom_preserves_each_environment_version(self):
        template = self.single_template()
        self.run_import(matrix=self.matrix_with_prom(), templates=[template], environment='IFT')
        prom_output = self.folder / 'prom_after_ift.xlsx'
        result = self.run_import(matrix=self.output, templates=[template], environment='PROM', output=prom_output)
        self.assertEqual((result.added, result.updated, result.versioned, result.promoted, result.prom_deactivated), (0, 1, 0, 1, 1))
        book = load_workbook(prom_output)
        sheet = book.worksheets[0]
        self.assertEqual((sheet['C2'].value, sheet['D2'].value), (None, None))
        self.assertEqual((sheet['C6'].value, sheet['D6'].value), (1, 1))
        self.assertEqual(sheet.max_row, 6)
        book.close()

    def test_other_environment_only_does_not_get_modified(self):
        template = self.single_template()
        for environment in ('IFT', 'PROM'):
            def edit(book):
                book.worksheets[0]['C2'] = None if environment == 'IFT' else 1
                book.worksheets[0]['D2'] = 1 if environment == 'IFT' else None
            matrix = self.change(self.matrix, edit, f'other_{environment}.xlsx')
            output = self.folder / f'isolated_{environment}.xlsx'
            result = self.run_import(matrix=matrix, templates=[template], environment=environment, output=output)
            self.assertEqual((result.added, result.updated, result.versioned), (1, 0, 0) if environment == 'IFT' else (0, 0, 0))
            self.assertEqual(len(result.conflicts), 0 if environment == 'IFT' else 1)
            before, after = load_workbook(matrix), load_workbook(output)
            self.assertEqual([c.value for c in before.worksheets[0][2]], [c.value for c in after.worksheets[0][2]])
            before.close()
            after.close()

    def test_inactive_history_does_not_conflict_with_active_version(self):
        def edit(book):
            sheet = book.worksheets[0]
            sheet.append([c.value for c in sheet[2]])
            sheet['C6'], sheet['D6'] = None, None
        matrix = self.change(self.matrix, edit, 'history.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()])
        self.assertEqual((result.added, result.updated, result.skipped), (0, 1, 0))
        self.assertFalse(result.conflicts)

    def test_formula_or_invalid_selected_flag_requires_manual_review(self):
        for value in ('=1', 'непонятно', True, 2):
            matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['C2'], 'value', value), f'flag_{str(value).replace("=", "")}.xlsx')
            result = self.run_import(matrix=matrix, templates=[self.single_template()], output=self.folder / f'flag_result_{str(value).replace("=", "")}.xlsx')
            self.assertEqual((result.added, result.updated), (0, 0))
            self.assertEqual(len(result.conflicts), 1)

    def test_text_flag_one_is_supported(self):
        matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['C2'], 'value', ' 1 '), 'text_flag.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()])
        self.assertEqual(result.updated, 1)

    def test_financial_type_matches_financial_only(self):
        matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['M2'], 'value', 'Финансовые операции'), 'financial.xlsx')
        template = self.single_template(lambda s: setattr(s['G18'], 'value', 'Финансовые операции'))
        result = self.run_import(matrix=matrix, templates=[template])
        self.assertEqual((result.added, result.updated), (0, 1))

    def test_blank_type_adds_nonfinancial_beside_financial(self):
        matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['M2'], 'value', 'Финансовые операции'), 'financial_old.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()])
        self.assertEqual((result.added, result.updated), (1, 0))
        book = load_workbook(self.output)
        sheet = book.worksheets[0]
        self.assertEqual(sheet['M2'].value, 'Финансовые операции')
        self.assertEqual(sheet['M6'].value, 'Не финансовые операции')
        self.assertEqual(sheet['C2'].value, 1)
        book.close()

    def test_blank_and_explicit_nonfinancial_are_same_key(self):
        explicit = self.single_template(lambda s: setattr(s['G18'], 'value', 'Не финансовые операции'))
        result = self.run_import(templates=[explicit])
        self.assertEqual((result.added, result.updated), (0, 1))

    def test_default_transaction_uses_actual_catalog_code(self):
        def edit(book):
            sheet = book['Типы транзакций']
            for row in sheet:
                if row[1].value == 'Не финансовые операции':
                    row[0].value = 'REAL_CODE_7'
        matrix = self.change(self.matrix, edit, 'real_code.xlsx')
        self.run_import(matrix=matrix)
        book = load_workbook(self.output)
        try:
            config = load_config()
            sheet = book.worksheets[0]
            record = read_matrix_row(sheet, columns_at(sheet, 1), 6, Catalog(book, config), config)
            self.assertEqual(record.values['TRANSACTION_TYPE'], 'REAL_CODE_7')
            self.assertEqual(sheet['M6'].value, 'Не финансовые операции')
        finally:
            book.close()

    def test_missing_default_transaction_is_not_guessed(self):
        def edit(book):
            sheet = book['Типы транзакций']
            for row in list(sheet):
                if row[1].value == 'Не финансовые операции':
                    sheet.delete_rows(row[0].row)
        matrix = self.change(self.matrix, edit, 'missing_type.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'Не финансовые операции'):
            self.run_import(matrix=matrix)
        self.assertFalse(self.output.exists())

    def test_bypass_formula_is_not_evaluated(self):
        result = self.run_import(templates=[self.single_template(lambda s: setattr(s['Q18'], 'value', '=1/0'))])
        self.assertEqual((result.added, result.updated), (0, 1))

    def test_legacy_defaults_cannot_enable_other_environment_or_bypass(self):
        config = load_config()
        config['new_row_defaults'].update(IFT=1, PROM=1, VALUE_SWITCH='ТЗ')
        path = self.folder / 'defaults.json'
        path.write_text(json.dumps(config, ensure_ascii=False), encoding='utf-8')
        self.run_import(config_path=path)
        book = load_workbook(self.output)
        sheet = book.worksheets[0]
        for row in range(6, 9):
            self.assertEqual(sheet.cell(row, 3).value, 1)
            self.assertIsNone(sheet.cell(row, 4).value)
            self.assertIsNone(sheet.cell(row, 38).value)
        book.close()

    def test_invalid_environment_is_rejected(self):
        with self.assertRaisesRegex(ImportProblem, 'среду задачи'):
            self.run_import(environment='BOTH')
        self.assertFalse(self.output.exists())

    def test_prom_unknown_ift_flag_is_manual(self):
        matrix = self.matrix_with_prom(ift='=1')
        result = self.run_import(matrix=matrix, templates=[self.single_template()], environment='PROM')
        self.assertEqual((result.added, result.updated), (0, 0))
        self.assertEqual(len(result.conflicts), 1)
        book = load_workbook(self.output)
        sheet = book.worksheets[0]
        self.assertEqual(sheet['C2'].value, '=1')
        self.assertEqual(sheet['D2'].value, 1)
        self.assertEqual(sheet.max_row, 5)
        book.close()

    def test_transaction_description_formula_is_not_assumed_blank(self):
        matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['M2'], 'value', '=1+1'), 'type_formula.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()])
        self.assertEqual((result.added, result.updated), (0, 0))
        self.assertEqual(len(result.conflicts), 1)

    def test_only_duplicates_active_for_selected_environment_conflict(self):
        def edit(book):
            sheet = book.worksheets[0]
            sheet['D2'] = 1
            sheet.append([c.value for c in sheet[2]])
            sheet['C6'] = None
        matrix = self.change(self.matrix, edit, 'env_duplicates.xlsx')
        template = self.single_template()
        ift = self.run_import(matrix=matrix, templates=[template], environment='IFT')
        self.assertEqual(ift.versioned, 1)
        self.assertFalse(ift.conflicts)
        prom = self.run_import(matrix=matrix, templates=[template], environment='PROM', output=self.folder / 'prom_duplicates.xlsx')
        self.assertEqual((prom.added, prom.updated), (0, 0))
        self.assertEqual(len(prom.conflicts), 1)

    def prepared_ift(self, old_prom=False, old_ift=None):
        def edit(book):
            sheet = book.worksheets[0]
            if old_prom:
                sheet['C2'], sheet['D2'] = old_ift, 1
                sheet.append([c.value for c in sheet[2]])
                sheet['C6'], sheet['D6'] = 1, None
                target = 6
            else:
                target = 2
            sheet.cell(target, 37, 'Новое тестовое сообщение. Повторите операцию позже.')
            sheet.cell(target, 38, 'ТЗ')
        return self.change(self.matrix, edit, 'prepared.xlsx')

    def assert_only_flags_changed(self, before_path, output_path, changes):
        before, after = load_workbook(before_path), load_workbook(output_path)
        try:
            a, b = before.worksheets[0], after.worksheets[0]
            self.assertEqual(a.max_row, b.max_row)
            for row in range(1, a.max_row + 1):
                for col in range(1, a.max_column + 1):
                    coord = a.cell(row, col).coordinate
                    expected = changes.get(coord, a.cell(row, col).value)
                    self.assertEqual(b.cell(row, col).value, expected, coord)
            for a, b in zip(before, after):
                if a is not before.worksheets[0]:
                    self.assertEqual(list(a.values), list(b.values))
        finally:
            before.close()
            after.close()

    def test_prom_activates_ready_ift_without_adding_or_editing_data(self):
        matrix = self.prepared_ift()
        result = self.run_import(matrix=matrix, templates=[self.single_template()], environment='PROM')
        self.assertEqual((result.added, result.updated, result.promoted, result.prom_deactivated), (0, 1, 1, 0))
        self.assertFalse(result.conflicts)
        self.assert_only_flags_changed(matrix, self.output, {'D2': 1})

    def test_prom_switches_from_old_row_to_ready_ift_preserving_both_ift_flags(self):
        matrix = self.prepared_ift(old_prom=True, old_ift=1)
        result = self.run_import(matrix=matrix, templates=[self.single_template()], environment='PROM')
        self.assertEqual((result.added, result.updated, result.promoted, result.prom_deactivated), (0, 1, 1, 1))
        self.assertFalse(result.conflicts)
        self.assert_only_flags_changed(matrix, self.output, {'D2': None, 'D6': 1})

    def test_prom_only_replacement_preserves_empty_ift_flag(self):
        matrix = self.prepared_ift(old_prom=True)
        result = self.run_import(matrix=matrix, templates=[self.single_template()], environment='PROM')
        self.assertEqual(result.prom_deactivated, 1)
        self.assert_only_flags_changed(matrix, self.output, {'D2': None, 'D6': 1})

    def test_promotion_repeat_is_idempotent(self):
        matrix = self.prepared_ift(old_prom=True)
        template = self.single_template()
        self.run_import(matrix=matrix, templates=[template], environment='PROM')
        repeat = self.folder / 'prom_repeat.xlsx'
        result = self.run_import(matrix=self.output, templates=[template], environment='PROM', output=repeat)
        self.assertEqual((result.added, result.updated, result.skipped, result.promoted, result.prom_deactivated), (0, 0, 1, 0, 0))
        self.assert_only_flags_changed(self.output, repeat, {})

    def test_duplicate_ready_ift_does_not_deactivate_old_prom(self):
        matrix = self.prepared_ift(old_prom=True)
        def edit(book):
            sheet = book.worksheets[0]
            sheet.append([c.value for c in sheet[6]])
        matrix = self.change(matrix, edit, 'duplicate_ready.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()], environment='PROM')
        self.assertEqual((result.added, result.updated), (0, 0))
        self.assertEqual(len(result.conflicts), 1)
        self.assert_only_flags_changed(matrix, self.output, {})

    def test_duplicate_active_prom_does_not_activate_ready_ift(self):
        matrix = self.prepared_ift(old_prom=True)
        def edit(book):
            sheet = book.worksheets[0]
            sheet.append([c.value for c in sheet[2]])
        matrix = self.change(matrix, edit, 'duplicate_prom.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()], environment='PROM')
        self.assertEqual((result.added, result.updated), (0, 0))
        self.assertEqual(len(result.conflicts), 1)
        self.assert_only_flags_changed(matrix, self.output, {})

    def test_mismatched_ift_payload_is_not_rewritten_for_promotion(self):
        for field in ('L18', 'O18', 'P18'):
            template = self.single_template(lambda s: setattr(s[field], 'value', 'З' if field == 'L18' else 'Другая версия'))
            matrix = self.prepared_ift(old_prom=True)
            output = self.folder / f'mismatch_{field}.xlsx'
            result = self.run_import(matrix=matrix, templates=[template], environment='PROM', output=output)
            self.assertEqual((result.added, result.updated), (0, 0))
            self.assertEqual(len(result.conflicts), 1)
            self.assert_only_flags_changed(matrix, output, {})

    def test_different_bypass_does_not_prevent_promotion(self):
        matrix = self.prepared_ift()
        template = self.single_template(lambda s: setattr(s['Q18'], 'value', '=1/0'))
        result = self.run_import(matrix=matrix, templates=[template], environment='PROM')
        self.assertEqual(result.promoted, 1)
        self.assert_only_flags_changed(matrix, self.output, {'D2': 1})

    def test_identical_template_duplicates_only_promote_once(self):
        matrix, template = self.prepared_ift(), self.single_template()
        result = self.run_import(matrix=matrix, templates=[template, template], environment='PROM')
        self.assertEqual((result.updated, result.skipped, result.promoted), (1, 1, 1))
        self.assert_only_flags_changed(matrix, self.output, {'D2': 1})

    def test_conflicting_templates_do_not_touch_prom_flags(self):
        matrix, template = self.prepared_ift(old_prom=True), self.single_template()
        other = self.change(template, lambda b: setattr(b.worksheets[0]['P18'], 'value', 'Противоречие'), 'conflicting_promotion.xlsx')
        result = self.run_import(matrix=matrix, templates=[template, other], environment='PROM')
        self.assertEqual((result.added, result.updated), (0, 0))
        self.assertEqual(len(result.conflicts), 2)
        self.assert_only_flags_changed(matrix, self.output, {})

    def test_prom_can_switch_to_ready_row_with_same_payload_as_old_prom(self):
        matrix = self.prepared_ift(old_prom=True)
        matrix = self.change(matrix, lambda b: setattr(b.worksheets[0]['AK2'], 'value', b.worksheets[0]['AK6'].value), 'same_prom_content.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()], environment='PROM')
        self.assertEqual((result.promoted, result.prom_deactivated), (1, 1))
        self.assert_only_flags_changed(matrix, self.output, {'D2': None, 'D6': 1})
