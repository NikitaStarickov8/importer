"""Нагрузочная проверка: python3 tests/benchmark.py --rows 50000."""
import argparse
import re
import sys
import tempfile
import time
import warnings
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'vendor'), str(ROOT)]
from openpyxl import Workbook, load_workbook
from app.importer import import_files


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--rows', type=int, default=50000)
    args = parser.parse_args()
    warnings.filterwarnings('ignore', message='Title is more than 31 characters')
    source = load_workbook(ROOT / 'examples/01_matrix.xlsx')
    base_sheet = source.worksheets[0]
    seeds = [[c.value for c in base_sheet[i]] for i in range(2, 6)]
    patterns = [re.sub(r'(?<![A-Z0-9_$])([A-Z]{1,3})2\b', r'\1{row}', v) if isinstance(v, str) and v.startswith('=') else v for v in seeds[0]]
    with tempfile.TemporaryDirectory(prefix='matrix-benchmark-') as folder:
        folder = Path(folder)
        matrix = folder / 'large.xlsx'
        output = folder / 'result.xlsx'
        write_start = time.monotonic()
        book = Workbook(write_only=True)
        sheet = book.create_sheet(base_sheet.title)
        sheet.append([c.value for c in base_sheet[1]])
        for i in range(args.rows):
            row = i + 2
            if i < 4:
                values = seeds[i]
            else:
                values = [v.format(row=row) if isinstance(v, str) and v.startswith('=') else v for v in patterns]
                values[0] = f'=ROW(A{row+1})-2'
                values[33] = i
            sheet.append(values)
        for original in source.worksheets[1:]:
            dest = book.create_sheet(original.title)
            for row in original.values:
                dest.append(row)
        book.save(matrix)
        source.close()
        print(f'Подготовлено {args.rows} строк за {time.monotonic()-write_start:.2f} с, размер {matrix.stat().st_size/1024/1024:.2f} МБ', flush=True)
        start = time.monotonic()
        result = import_files(matrix, [ROOT/'examples/02_template_update.xlsx', ROOT/'examples/03_template_new_group.xlsx'], output,
                              progress=lambda s: print(s, flush=True))
        elapsed = time.monotonic()-start
        assert (result.added, result.updated, result.skipped) == (3, 1, 1), result
        check = load_workbook(output, read_only=True)
        assert check.worksheets[0].max_row == args.rows + 4
        check.close()
        print(f'Импорт: {elapsed:.2f} с; +{result.added}, обновлено {result.updated}, без изменений {result.skipped}', flush=True)


if __name__ == '__main__':
    main()
