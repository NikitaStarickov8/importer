"""Проверки на копиях искусственных Excel-файлов; originals не изменяются."""
import hashlib
import json
import sys
import tempfile
import unittest
import warnings
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "vendor"), str(ROOT)]
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from app.importer import Catalog, ImportProblem, import_files, load_config

warnings.filterwarnings("ignore", message="Title is more than 31 characters")


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.matrix = ROOT / "examples/01_matrix.xlsx"
        self.templates = [ROOT / "examples/02_template_update.xlsx", ROOT / "examples/03_template_new_group.xlsx"]
        self.output = self.folder / "result.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def change(self, path, edit, name):
        book = load_workbook(path)
        edit(book)
        out = self.folder / name
        book.save(out)
        book.close()
        return out

    def run_import(self, **kwargs):
        return import_files(kwargs.get("matrix", self.matrix), kwargs.get("templates", self.templates),
                            kwargs.get("output", self.output), kwargs.get("config_path"))

    def test_demo_and_repeat_without_cached_formula_values(self):
        hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in [self.matrix, *self.templates]}
        result = self.run_import()
        self.assertEqual((result.added, result.updated, result.skipped, result.total), (3, 1, 1, 5))
        repeat = self.run_import(matrix=self.output, output=self.folder / "repeat.xlsx")
        self.assertEqual((repeat.added, repeat.updated, repeat.skipped), (0, 0, 5))
        for p, expected in hashes.items():
            self.assertEqual(hashlib.sha256(p.read_bytes()).digest(), expected)

    def test_formulas_flags_unmentioned_rows_and_table(self):
        self.run_import()
        before, after = load_workbook(self.matrix), load_workbook(self.output)
        try:
            a, b = before.worksheets[0], after.worksheets[0]
            for row in range(2, 6):
                for col in (1, 3, 4, 6, 8, 10, 21, 27, 30, 32, 33, 35, 39, 40, 41, 42):
                    self.assertEqual(a.cell(row, col).value, b.cell(row, col).value)
            for row in (4, 5):
                self.assertEqual([c.value for c in a[row]], [c.value for c in b[row]])
            self.assertEqual(b['AK2'].value, 'Новое тестовое сообщение. Повторите операцию позже.')
            self.assertEqual(b['AL6'].value, 'ТЗ')
            self.assertIsNone(b['AL8'].value)
            self.assertEqual(b['A6'].value, '=ROW(A7)-2')
            self.assertIn('MATCH(G6,', b['H6'].value)
            self.assertIsNone(b['C6'].value)
            self.assertIsNone(b['D6'].value)
            self.assertIsNone(b['B6'].value)
            self.assertEqual(b.tables['MatrixTable'].ref, 'A1:AQ8')
            self.assertTrue(after.calculation.fullCalcOnLoad)
            self.assertTrue(after.calculation.forceFullCalc)
            self.assertEqual(before.sheetnames, after.sheetnames)
            for source, dest in zip(before.worksheets[1:], after.worksheets[1:]):
                self.assertEqual(list(source.values), list(dest.values))
        finally:
            before.close()
            after.close()

    def test_unknown_lookup_stops_entire_batch(self):
        bad = self.change(self.templates[1], lambda b: setattr(b.worksheets[0]['J18'], 'value', 'UNKNOWN STATUS'), 'bad.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'UNKNOWN STATUS'):
            self.run_import(templates=[self.templates[0], bad])
        self.assertFalse(self.output.exists())

    def test_duplicate_target_is_not_updated(self):
        def duplicate(book):
            s = book.worksheets[0]
            s.append([c.value for c in s[2]])
        bad = self.change(self.matrix, duplicate, 'duplicates.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'несколько подходящих'):
            self.run_import(matrix=bad)
        self.assertFalse(self.output.exists())

    def test_duplicate_templates_are_skipped(self):
        result = self.run_import(templates=[*self.templates, self.templates[0]])
        self.assertEqual((result.added, result.updated, result.skipped), (3, 1, 4))

    def test_conflicting_templates_stop_before_save(self):
        bad = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['P18'], 'value', 'Конфликтующее сообщение'), 'conflict.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'Конфликт шаблонов'):
            self.run_import(templates=[self.templates[0], bad])
        self.assertFalse(self.output.exists())

    def test_unknown_bypass(self):
        bad = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['Q18'], 'value', 'неизвестно'), 'bypass.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'Байпас'):
            self.run_import(templates=[bad])

    def test_bypass_clears_previous_forbidden_value(self):
        matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['AL2'], 'value', 'ТЗ'), 'bypass_matrix.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.templates[0]])
        self.assertEqual(result.updated, 1)
        b = load_workbook(self.output)
        self.assertIsNone(b.worksheets[0]['AL2'].value)
        b.close()

    def test_required_field(self):
        bad = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['J18'], 'value', None), 'empty.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'STOP_LIST_STATUS'):
            self.run_import(templates=[bad])

    def test_unknown_formula_not_read_as_blank(self):
        bad = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['P18'], 'value', '=1+1'), 'formula.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'P18.*формула'):
            self.run_import(templates=[bad])

    def test_literal_equals_text_remains_text(self):
        def edit(book):
            c = book.worksheets[0]['P18']
            c.value = '=это текст'
            c.data_type = 's'
        template = self.change(self.templates[0], edit, 'text.xlsx')
        self.run_import(templates=[template])
        b = load_workbook(self.output)
        self.assertEqual(b.worksheets[0]['AK2'].value, '=это текст')
        self.assertEqual(b.worksheets[0]['AK2'].data_type, 's')
        b.close()

    def test_zero_amount_is_different_from_empty(self):
        template = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['N18'], 'value', 0), 'zero.xlsx')
        r = self.run_import(templates=[template])
        self.assertEqual((r.added, r.updated, r.skipped), (2, 0, 1))
        b = load_workbook(self.output)
        self.assertEqual(b.worksheets[0]['AH6'].value, 0)
        b.close()

    def test_negative_amount(self):
        bad = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['N18'], 'value', -1), 'negative.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'неотрицательным'):
            self.run_import(templates=[bad])

    def test_platform_changes_make_new_group_for_sbol_pro(self):
        def edit(book):
            for row in range(18, 21):
                book.worksheets[0].cell(row, 5, 'Web — тестовая платформа')
        template = self.change(self.templates[0], edit, 'platform.xlsx')
        r = self.run_import(templates=[template])
        self.assertEqual((r.added, r.updated), (3, 0))

    def test_non_sbol_platform_is_not_in_key(self):
        def matrix_edit(book):
            for row in range(2, 6):
                book.worksheets[0].cell(row, 5, 'Офис')
        def template_edit(book):
            for row in range(18, 21):
                book.worksheets[0].cell(row, 1, 'Офис')
                book.worksheets[0].cell(row, 5, 'Web — тестовая платформа')
        matrix = self.change(self.matrix, matrix_edit, 'office.xlsx')
        template = self.change(self.templates[0], template_edit, 'office_template.xlsx')
        r = self.run_import(matrix=matrix, templates=[template])
        self.assertEqual((r.added, r.updated, r.skipped), (1, 2, 0))

    def test_input_and_existing_output_never_overwritten(self):
        before = self.matrix.read_bytes()
        with self.assertRaises(ImportProblem):
            self.run_import(output=self.matrix)
        self.assertEqual(self.matrix.read_bytes(), before)
        self.output.write_bytes(b'existing content')
        with self.assertRaises(ImportProblem):
            self.run_import()
        self.assertEqual(self.output.read_bytes(), b'existing content')

    def test_duplicate_catalog_descriptions_rejected(self):
        bad = self.change(self.matrix, lambda b: b['Каналы'].append(['SECOND_CODE', 'СБОЛ про']), 'bad_catalog.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'нескольким записям'):
            self.run_import(matrix=bad)

    def test_extend_validations_and_conditional_formatting(self):
        def edit(book):
            s = book.worksheets[0]
            dv = DataValidation(type='list', formula1='"Р,З,ТЗ"')
            s.add_data_validation(dv)
            dv.add('AE2:AE5')
            s.conditional_formatting.add('AH2:AH5', CellIsRule(operator='greaterThan', formula=['0']))
        matrix = self.change(self.matrix, edit, 'validation.xlsx')
        self.run_import(matrix=matrix)
        b = load_workbook(self.output)
        self.assertEqual(str(b.worksheets[0].data_validations.dataValidation[0].sqref), 'AE2:AE8')
        self.assertEqual(str(next(iter(b.worksheets[0].conditional_formatting)).sqref), 'AH2:AH8')
        b.close()

    def test_reordered_template_columns(self):
        def edit(book):
            s = book.worksheets[0]
            s.insert_cols(12)
        template = self.change(self.templates[0], edit, 'reordered.xlsx')
        r = self.run_import(templates=[template])
        self.assertEqual((r.added, r.updated, r.skipped), (1, 1, 1))

    def test_prefilled_table_spare_row_grows_to_new_data_end(self):
        def edit(book):
            book.worksheets[0].tables['MatrixTable'].ref = 'A1:AQ6'
        matrix = self.change(self.matrix, edit, 'spare.xlsx')
        self.run_import(matrix=matrix)
        b = load_workbook(self.output)
        self.assertEqual(b.worksheets[0].tables['MatrixTable'].ref, 'A1:AQ8')
        b.close()

    def test_data_after_table_cannot_be_overwritten(self):
        matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['AK6'], 'value', 'Заметка сотрудника'), 'notes.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'содержит данные'):
            self.run_import(matrix=matrix)
        self.assertFalse(self.output.exists())

    def test_top_d3_to_d6_are_simple_links_and_not_needed_by_complete_rows(self):
        def edit(book):
            sheet = book.worksheets[0]
            for row in range(3, 7):
                sheet[f'D{row}'] = f'=B{row}'
        template = self.change(self.templates[0], edit, 'top_links.xlsx')
        result = self.run_import(templates=[template])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_unneeded_complex_header_formula_is_not_evaluated(self):
        def edit(book):
            sheet = book.worksheets[0]
            for address in ['D3','D4','D5','D6','D7','B8','D10']:
                sheet[address] = '=UNSUPPORTED_FUNCTION(123)'
        template = self.change(self.templates[0], edit, 'unused_header.xlsx')
        result = self.run_import(templates=[template])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_row_references_follow_header_links_without_cached_results(self):
        def edit(book):
            sheet = book.worksheets[0]
            for row in range(3, 8):
                sheet[f'D{row}'] = f'=$B${row}'
            # В строке передаём коды из D5/D7; остальные описания идут из шапки.
            for row in range(18, 21):
                for column, address in [('A','D3'),('B','D5'),('C','D7'),('D','B10'),('H','B8')]:
                    sheet[f'{column}{row}'] = f'=${address[0]}${address[1:]}'
        template = self.change(self.templates[0], edit, 'row_links.xlsx')
        result = self.run_import(templates=[template])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_missing_row_values_are_taken_from_linked_header(self):
        def edit(book):
            sheet = book.worksheets[0]
            for row in range(3, 8):
                sheet[f'D{row}'] = f'=B{row}'
            for row in range(18, 21):
                for col in (1, 2, 3, 4, 8):
                    sheet.cell(row, col).value = None
        template = self.change(self.templates[0], edit, 'header_fallback.xlsx')
        result = self.run_import(templates=[template])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_quoted_sheet_reference_and_apostrophe(self):
        def edit(book):
            sheet = book.worksheets[0]
            helper = book.create_sheet("Общие 'поля'")
            helper['B3'] = 'СБОЛ про'
            sheet['D3'] = "='Общие ''поля'''!$B$3"
            for row in range(18, 21):
                sheet[f'A{row}'] = '=$D$3'
        template = self.change(self.templates[0], edit, 'quoted_links.xlsx')
        result = self.run_import(templates=[template])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_formula_link_cycle_is_rejected(self):
        def edit(book):
            sheet = book.worksheets[0]
            sheet['D3'], sheet['B3'] = '=B3', '=D3'
            sheet['A18'] = '=D3'
        template = self.change(self.templates[0], edit, 'cycle.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'циклическая ссылка'):
            self.run_import(templates=[template])
        self.assertFalse(self.output.exists())

    def test_blank_link_can_fall_back_to_common_value(self):
        def edit(book):
            sheet = book.worksheets[0]
            sheet['A18'] = '=$F$2'
            sheet['D3'] = '=B3'
        template = self.change(self.templates[0], edit, 'blank_link.xlsx')
        result = self.run_import(templates=[template])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_linked_text_preserves_newlines_and_equals_literal(self):
        def edit(book):
            sheet = book.worksheets[0]
            sheet['F2'] = '=обычный текст\nВторая строка'
            sheet['F2'].data_type = 's'
            sheet['P18'] = '=$F$2'
        template = self.change(self.templates[0], edit, 'linked_text.xlsx')
        self.run_import(templates=[template])
        book = load_workbook(self.output)
        self.assertEqual(book.worksheets[0]['AK2'].value, '=обычный текст\nВторая строка')
        self.assertEqual(book.worksheets[0]['AK2'].data_type, 's')
        book.close()

    def test_unrelated_incomplete_client_rows_do_not_block_import(self):
        def edit(book):
            sheet = book['Типы Клиентов']
            sheet.append(['UNUSED', 20, None])
            sheet.append([None, 30, 'Другой незаполненный тип'])
            sheet.append([None, 40, None])
        matrix = self.change(self.matrix, edit, 'incomplete_catalog.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.templates[0]])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))
        book = load_workbook(self.output)
        self.assertEqual(book['Типы Клиентов']['A3'].value, 'UNUSED')
        self.assertIsNone(book['Типы Клиентов']['C3'].value)
        self.assertEqual(book['Типы Клиентов']['C4'].value, 'Другой незаполненный тип')
        book.close()

    def test_simple_links_in_catalog_are_resolved(self):
        def edit(book):
            sheet = book['Типы Клиентов']
            helper = book.create_sheet('Значения справочника')
            helper['A1'], helper['B1'] = 'PRODUCT_OWNER', 'Владелец продукта'
            sheet['A2'] = "='Значения справочника'!$A$1"
            sheet['C2'] = "='Значения справочника'!$B$1"
        matrix = self.change(self.matrix, edit, 'catalog_links.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.templates[0]])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_requested_incomplete_client_entry_reports_exact_cell(self):
        matrix = self.change(self.matrix, lambda b: setattr(b['Типы Клиентов']['A2'], 'value', None), 'missing_client_code.xlsx')
        with self.assertRaises(ImportProblem) as context:
            self.run_import(matrix=matrix, templates=[self.templates[0]])
        message = str(context.exception)
        self.assertIn('Типы Клиентов', message)
        self.assertIn('Владелец продукта', message)
        self.assertIn('строка 2', message)
        self.assertIn('A2', message)
        self.assertFalse(self.output.exists())

    def test_lookup_by_code_does_not_bypass_missing_description(self):
        matrix = self.change(self.matrix, lambda b: setattr(b['Типы Клиентов']['C2'], 'value', None), 'missing_client_description.xlsx')
        template = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['H18'], 'value', 'PRODUCT_OWNER'), 'client_by_code.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'C2.*пустое поле'):
            self.run_import(matrix=matrix, templates=[template])
        self.assertFalse(self.output.exists())

    def test_requested_catalog_formula_reports_its_cell(self):
        matrix = self.change(self.matrix, lambda b: setattr(b['Типы Клиентов']['A2'], 'value', '=1+1'), 'catalog_formula.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'A2.*неожиданная формула'):
            self.run_import(matrix=matrix, templates=[self.templates[0]])
        self.assertFalse(self.output.exists())

    def test_unrelated_formula_catalog_row_does_not_block_literal_entry(self):
        matrix = self.change(self.matrix, lambda b: b['Типы Клиентов'].append(['=1+1', 20, 'Неиспользуемый тип']), 'unused_formula.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.templates[0]])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_missing_value_diagnostics_include_incomplete_catalog_rows(self):
        matrix = self.change(self.matrix, lambda b: setattr(b['Типы Клиентов']['C2'], 'value', None), 'missing_description.xlsx')
        with self.assertRaises(ImportProblem) as context:
            self.run_import(matrix=matrix, templates=[self.templates[0]])
        self.assertIn('не найдено', str(context.exception))
        self.assertIn('C2', str(context.exception))
        self.assertIn('Типы Клиентов', str(context.exception))
        self.assertFalse(self.output.exists())

    def test_bad_duplicate_alias_is_not_ignored(self):
        matrix = self.change(self.matrix, lambda b: b['Типы Клиентов'].append([None, 20, 'Владелец продукта']), 'bad_duplicate.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'A3'):
            self.run_import(matrix=matrix, templates=[self.templates[0]])
        self.assertFalse(self.output.exists())

    def test_template_filename_is_not_duplicated_in_error(self):
        template = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['J18'], 'value', 'UNKNOWN'), 'unique_template_name.xlsx')
        with self.assertRaises(ImportProblem) as context:
            self.run_import(templates=[template])
        self.assertEqual(str(context.exception).count('unique_template_name.xlsx'), 1)

    def test_template_needs_only_its_three_reference_sheets(self):
        def edit(book):
            sheet = book.worksheets[0]
            book.remove(book['Типы Клиентов'])
            for title in ('Каналы', 'COMPLIANCE_STOP_LIST_CODES', 'Статусы Стоп Листа'):
                reference = book.create_sheet(title)
                reference.append(['Код', 'Описание'])
                reference.append(['WRONG_TEMPLATE_CODE', 'СБОЛ про'])
            # Клиент берётся из B8 и справочника основной матрицы, а не из D8.
            sheet['H18'] = None
        template = self.change(self.templates[0], edit, 'three_reference_sheets.xlsx')
        result = self.run_import(templates=[template])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_catalog_error_names_main_matrix_as_source(self):
        matrix = self.change(self.matrix, lambda b: setattr(b['Типы Клиентов']['A2'], 'value', None), 'main_matrix_bad_catalog.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'основной матрице.*Типы Клиентов'):
            self.run_import(matrix=matrix, templates=[self.templates[0]])

    def marker_files(self, label='Флаг 66'):
        def edit_matrix(book):
            sheet = book['Статусы Стоп Листа']
            sheet['A5'], sheet['B5'], sheet['C5'] = 'MARKER_66', 'Маркер 66', 'NORMAL'
            sheet = book['COMPLIANCE_STOP_LIST_CODES']
            sheet['A5'], sheet['B5'] = 'MARKER_66', 'Маркер 66'
        matrix = self.change(self.matrix, edit_matrix, 'marker_matrix.xlsx')
        template = self.change(self.templates[0], lambda b: setattr(b.worksheets[0]['J20'], 'value', label), 'marker_template.xlsx')
        return matrix, template

    def alias_config(self, aliases):
        config = load_config()
        config['aliases'] = aliases
        path = self.folder / 'aliases.json'
        path.write_text(json.dumps(config, ensure_ascii=False), encoding='utf-8')
        return path

    def test_flag_66_alias_writes_actual_description_and_is_idempotent(self):
        matrix, template = self.marker_files()
        result = self.run_import(matrix=matrix, templates=[template])
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))
        book = load_workbook(self.output)
        self.assertEqual(book.worksheets[0]['Z6'].value, 'Маркер 66')
        self.assertIn('MATCH(Z6,', book.worksheets[0]['AA6'].value)
        self.assertEqual(book['Статусы Стоп Листа']['A5'].value, 'MARKER_66')
        book.close()
        again = self.run_import(matrix=self.output, templates=[template], output=self.folder/'again.xlsx')
        self.assertEqual((again.added, again.updated, again.skipped), (0, 0, 3))

    def test_alias_matches_case_and_nonbreaking_spaces(self):
        matrix, template = self.marker_files('  ФЛАГ\xa066  ')
        result = self.run_import(matrix=matrix, templates=[template])
        self.assertEqual(result.added, 1)

    def test_alias_is_scoped_to_lookup_field(self):
        matrix, _ = self.marker_files()
        book = load_workbook(matrix)
        try:
            book['Подсистемы']['B2'] = 'Флаг 66'
            catalog = Catalog(book, load_config())
            self.assertEqual(catalog.resolve('SP_NAME', 'Флаг 66').code, 'DEMO_LOAN')
            self.assertEqual(catalog.resolve('STOP_LIST_STATUS', 'Флаг 66').code, 'MARKER_66')
        finally:
            book.close()

    def test_alias_does_not_replace_client_message_text(self):
        matrix, template = self.marker_files()
        template = self.change(template, lambda b: setattr(b.worksheets[0]['P20'], 'value', 'Проверьте Флаг 66'), 'literal_alias_text.xlsx')
        self.run_import(matrix=matrix, templates=[template])
        book = load_workbook(self.output)
        self.assertEqual(book.worksheets[0]['AK6'].value, 'Проверьте Флаг 66')
        book.close()

    def test_alias_missing_target_is_an_explicit_configuration_error(self):
        matrix, template = self.marker_files()
        config = self.alias_config({'STOP_LIST_STATUS':{'Флаг 66':'MISSING_66'}})
        with self.assertRaisesRegex(ImportProblem, 'aliases.STOP_LIST_STATUS.*MISSING_66'):
            self.run_import(matrix=matrix, templates=[template], config_path=config)
        self.assertFalse(self.output.exists())

    def test_alias_target_must_be_actual_code(self):
        matrix, template = self.marker_files()
        config = self.alias_config({'STOP_LIST_STATUS':{'Флаг 66':'Маркер 66'}})
        with self.assertRaisesRegex(ImportProblem, 'технический код'):
            self.run_import(matrix=matrix, templates=[template], config_path=config)

    def test_alias_cannot_override_existing_different_lookup(self):
        matrix, template = self.marker_files()
        config = self.alias_config({'STOP_LIST_STATUS':{'В стоп-листе не найден':'MARKER_66'}})
        with self.assertRaisesRegex(ImportProblem, 'конфликтует'):
            self.run_import(matrix=matrix, templates=[template], config_path=config)

    def test_alias_cannot_hide_ambiguous_target_code(self):
        matrix, template = self.marker_files()
        matrix = self.change(matrix, lambda b: b['Статусы Стоп Листа'].append(['MARKER_66','Другое описание','NORMAL']), 'ambiguous_marker.xlsx')
        with self.assertRaisesRegex(ImportProblem, 'нескольким записям'):
            self.run_import(matrix=matrix, templates=[template])
        self.assertFalse(self.output.exists())

    def test_alias_setting_is_optional_for_older_configs(self):
        config = load_config()
        del config['aliases']
        path = self.folder/'old_config.json'
        path.write_text(json.dumps(config, ensure_ascii=False), encoding='utf-8')
        result = self.run_import(templates=[self.templates[0]], config_path=path)
        self.assertEqual((result.added, result.updated, result.skipped), (1, 1, 1))

    def test_invalid_alias_config_is_rejected_before_import(self):
        cases = [None, {'STOP_LIST_STATUS':['Флаг 66','MARKER_66']}, {'UNKNOWN_FIELD':{}},
                 {'STOP_LIST_STATUS':{'Флаг 66':''}},
                 {'STOP_LIST_STATUS':{'Флаг 66':'MARKER_66','ФЛАГ 66':'OTHER'}}]
        for aliases in cases:
            with self.subTest(aliases=aliases):
                config = self.alias_config(aliases)
                with self.assertRaisesRegex(ImportProblem, 'aliases'):
                    self.run_import(config_path=config)
        self.assertFalse(self.output.exists())


if __name__ == '__main__':
    unittest.main(verbosity=2)
