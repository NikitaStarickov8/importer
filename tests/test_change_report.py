"""Отчёт содержит только реально затронутые строки Excel, без выдуманного CODE."""
import json
import unittest
from dataclasses import asdict

import test_workflow as workflow_tests
from openpyxl import load_workbook


class ChangeReportTests(unittest.TestCase):
    setUp = workflow_tests.WorkflowTests.setUp
    tearDown = workflow_tests.WorkflowTests.tearDown
    change = workflow_tests.WorkflowTests.change
    run_import = workflow_tests.WorkflowTests.run_import
    single_template = workflow_tests.WorkflowTests.single_template
    matrix_with_prom = workflow_tests.WorkflowTests.matrix_with_prom
    prepared_ift = workflow_tests.WorkflowTests.prepared_ift

    def test_ift_report_lists_actual_added_and_updated_excel_rows(self):
        result = self.run_import()
        self.assertEqual(result.affected_rows, [2, 6, 7, 8])
        self.assertEqual([(c.excel_row, c.action) for c in result.changes], [(2, 'updated'), (6, 'added'), (7, 'added'), (8, 'added')])
        self.assertEqual(result.changes[0].changed_fields, ['TEXT_FOR_CLIENT'])
        self.assertIn('Затронуты строки Excel: 2, 6, 7, 8', result.change_report)
        self.assertEqual(len(result.outcomes), result.total)
        unchanged = [outcome for outcome in result.outcomes if outcome.status == 'unchanged']
        self.assertEqual(len(unchanged), 1)
        self.assertEqual(unchanged[0].excel_rows, [3])
        self.assertIn('!19', unchanged[0].source)
        self.assertIn('совпадают', unchanged[0].reason)
        self.assertIn('БЕЗ ИЗМЕНЕНИЙ', result.change_report)
        before, after = load_workbook(self.matrix), load_workbook(self.output)
        try:
            a, b = before.worksheets[0], after.worksheets[0]
            actual = [row for row in range(2, b.max_row + 1) if [c.value for c in a[row]] != [c.value for c in b[row]]]
            self.assertEqual(actual, result.affected_rows)
        finally:
            before.close()
            after.close()

    def test_version_report_links_old_and_new_rows(self):
        result = self.run_import(matrix=self.matrix_with_prom(), templates=[self.single_template()])
        self.assertEqual(result.affected_rows, [2, 6])
        self.assertEqual([(c.excel_row, c.action, c.related_row) for c in result.changes], [(2, 'ift_deactivated', 6), (6, 'version_added', 2)])
        self.assertEqual(result.changes[0].changed_fields, ['IFT'])
        self.assertIn('прежняя строка: 2', result.change_report)
        self.assertIn('новая версия: 6', result.change_report)

    def test_prom_report_lists_both_flag_changes(self):
        result = self.run_import(matrix=self.prepared_ift(old_prom=True, old_ift=1), templates=[self.single_template()], environment='PROM')
        self.assertEqual(result.affected_rows, [2, 6])
        self.assertEqual([(c.excel_row, c.action, c.related_row) for c in result.changes], [(2, 'prom_deactivated', 6), (6, 'prom_activated', 2)])
        self.assertTrue(all(c.changed_fields == ['PROM'] for c in result.changes))
        self.assertIn('снята отметка PROM', result.change_report)
        self.assertIn('поставлена отметка PROM', result.change_report)

    def test_prom_first_activation_reports_only_target(self):
        result = self.run_import(matrix=self.prepared_ift(), templates=[self.single_template()], environment='PROM')
        self.assertEqual(result.affected_rows, [2])
        self.assertEqual(result.changes[0].action, 'prom_activated')
        self.assertIsNone(result.changes[0].related_row)

    def test_repeat_has_empty_change_list(self):
        self.run_import()
        result = self.run_import(matrix=self.output, output=self.folder / 'repeat.xlsx')
        self.assertEqual(result.affected_rows, [])
        self.assertEqual(result.changes, [])
        self.assertIn('Изменённых строк нет.', result.change_report)
        self.assertEqual(len(result.outcomes), 5)
        self.assertTrue(all(outcome.status == 'unchanged' for outcome in result.outcomes))

    def test_conflicting_and_unchanged_rows_are_not_marked_as_changed(self):
        def edit(book):
            sheet = book.worksheets[0]
            sheet.append([c.value for c in sheet[2]])
        matrix = self.change(self.matrix, edit, 'duplicate.xlsx')
        result = self.run_import(matrix=matrix)
        self.assertEqual(len(result.conflicts), 1)
        self.assertEqual(result.affected_rows, [7, 8, 9])
        self.assertNotIn(2, result.affected_rows)
        self.assertNotIn(3, result.affected_rows)
        self.assertNotIn(6, result.affected_rows)
        conflicts = [outcome for outcome in result.outcomes if outcome.status == 'conflict']
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(conflicts[0].excel_rows, [2, 6])
        self.assertIn('несколько подходящих', conflicts[0].reason)
        self.assertIn('ПРОПУЩЕНА: РУЧНОЙ РАЗБОР', result.change_report)

    def test_duplicate_templates_do_not_duplicate_report_entries(self):
        template = self.single_template()
        result = self.run_import(templates=[template, template])
        self.assertEqual(result.affected_rows, [2])
        self.assertEqual(len(result.changes), 1)
        self.assertEqual([outcome.status for outcome in result.outcomes], ['updated', 'skipped'])
        self.assertEqual(result.outcomes[1].excel_rows, [2])
        self.assertIn('Первый источник', result.outcomes[1].reason)
        self.assertIn('ПРОПУЩЕНА: ПОВТОР', result.change_report)

    def test_formula_first_column_is_not_evaluated_or_used_as_row_number(self):
        result = self.run_import(templates=[self.single_template()])
        change = result.changes[0]
        self.assertEqual(change.excel_row, 2)
        self.assertEqual(change.first_column_name, 'CODE')
        self.assertEqual(change.first_column_kind, 'formula')
        self.assertIsNone(change.first_column_value)
        self.assertIn('значение после пересчёта формулы', result.change_report)
        self.assertNotIn('CODE=2', result.change_report)

    def test_literal_first_column_preserves_identifier(self):
        matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['A2'], 'value', '000123141'), 'literal_code.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()])
        change = result.changes[0]
        self.assertEqual(change.excel_row, 2)
        self.assertEqual(change.first_column_value, '000123141')
        self.assertEqual(change.first_column_kind, 'value')
        self.assertIn('CODE=000123141', result.change_report)

    def test_first_column_label_follows_actual_first_column(self):
        def edit(book):
            sheet = book.worksheets[0]
            for row in sheet:
                a, b = row[0].value, row[1].value
                row[0].value, row[1].value = b, a
        matrix = self.change(self.matrix, edit, 'reordered.xlsx')
        result = self.run_import(matrix=matrix, templates=[self.single_template()])
        self.assertEqual(result.changes[0].first_column_name, 'Столбец1')
        self.assertEqual(result.changes[0].first_column_value, 'DEMO-001')
        self.assertEqual(result.affected_rows, [2])

    def test_blank_and_error_identifiers_do_not_break_report(self):
        for value, kind in ((None, 'blank'), ('#N/A', 'error')):
            matrix = self.change(self.matrix, lambda b: setattr(b.worksheets[0]['A2'], 'value', value), f'{kind}.xlsx')
            result = self.run_import(matrix=matrix, templates=[self.single_template()], output=self.folder / f'{kind}_result.xlsx')
            self.assertEqual(result.changes[0].first_column_kind, kind)
            self.assertEqual(result.affected_rows, [2])

    def test_report_is_json_serializable_for_cli(self):
        result = self.run_import()
        data = json.loads(json.dumps(asdict(result), ensure_ascii=False))
        self.assertEqual(data['affected_rows'], [2, 6, 7, 8])
        self.assertEqual(data['changes'][0]['excel_row'], 2)
        self.assertEqual(data['change_report'], result.change_report)
        self.assertEqual(len(data['outcomes']), result.total)

    def test_ambiguous_template_has_an_outcome_for_every_skipped_row(self):
        matrix = self.change(self.matrix, lambda b: b['Каналы'].append(['OTHER_CODE', 'СБОЛ про']), 'ambiguous.xlsx')
        result = self.run_import(matrix=matrix)
        self.assertEqual(len(result.outcomes), result.total)
        self.assertEqual(len(result.outcomes), 5)
        self.assertTrue(all(outcome.status == 'conflict' for outcome in result.outcomes))
        self.assertTrue(all('нескольким записям' in outcome.reason for outcome in result.outcomes))
        self.assertEqual(result.affected_rows, [])

    def test_prom_missing_ift_has_reason_and_source_without_false_row(self):
        result = self.run_import(matrix=self.matrix_with_prom(ift=None), templates=[self.single_template()], environment='PROM')
        self.assertEqual(len(result.outcomes), 1)
        outcome = result.outcomes[0]
        self.assertEqual(outcome.status, 'conflict')
        self.assertEqual(outcome.excel_rows, [2])
        self.assertIn('нет активной ИФТ', outcome.reason)
        self.assertIn('!18', outcome.source)

    def test_prom_mismatch_reason_names_fields(self):
        matrix = self.prepared_ift()
        result = self.run_import(matrix=matrix, templates=[self.single_template(lambda s: setattr(s['P18'], 'value', 'Другой текст'))], environment='PROM')
        self.assertEqual(result.outcomes[0].status, 'conflict')
        self.assertIn('TEXT_FOR_CLIENT', result.outcomes[0].reason)
        self.assertEqual(result.outcomes[0].excel_rows, [2])

    def test_prom_already_deployed_explains_unchanged_result(self):
        template = self.single_template()
        self.run_import(matrix=self.prepared_ift(), templates=[template], environment='PROM')
        result = self.run_import(matrix=self.output, templates=[template], environment='PROM', output=self.folder / 'repeat_prom.xlsx')
        self.assertEqual(len(result.outcomes), 1)
        self.assertEqual(result.outcomes[0].status, 'unchanged')
        self.assertIn('уже действуют в ПРОМ', result.outcomes[0].reason)
        self.assertEqual(result.outcomes[0].excel_rows, [2])
