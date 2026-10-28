"""Импорт без запуска Excel. Все изменения сохраняются только после проверки пакета."""
from __future__ import annotations

import json
import os
import re
import tempfile
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

ROOT = Path(__file__).resolve().parent.parent
MUTABLE_FIELDS = ("VALUE", "TITLE_FOR_CLIENT", "TEXT_FOR_CLIENT")
DEFAULT_TRANSACTION_DESCRIPTION = "Не финансовые операции"


class ImportProblem(Exception):
    """Понятная пользователю ошибка; выходной файл не создаётся."""


class AmbiguousLookup(ImportProblem):
    """Несколько соответствий: строка требует ручного разбора."""


def norm(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip().casefold()


def header(value):
    return re.sub(r"\s+", "", str(value or "").replace("\xa0", "")).upper()


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def transaction_default(config):
    value = config.get("transaction_type_default_description", DEFAULT_TRANSACTION_DESCRIPTION)
    if not isinstance(value, str) or blank(value):
        raise ImportProblem("transaction_type_default_description: требуется непустое описание типа транзакции.")
    return value


def validate_environment(environment):
    if environment not in ("IFT", "PROM"):
        raise ImportProblem("Выберите среду задачи: ИФТ (IFT) или ПРОМ (PROM).")
    return environment


def formula(value):
    return isinstance(value, str) and value.startswith("=")


def configured_aliases(config):
    aliases = config.get("aliases", {})
    if not isinstance(aliases, dict):
        raise ImportProblem("Настройка aliases должна быть объектом с именами полей.")
    result = {}
    for name, mapping in aliases.items():
        if name not in config["lookups"]:
            raise ImportProblem(f"aliases: неизвестное поле {name}. Используйте имя из lookups, например STOP_LIST_STATUS.")
        if not isinstance(mapping, dict):
            raise ImportProblem(f"aliases.{name}: требуется объект «название в шаблоне»: «код в матрице».")
        values = {}
        for source, target in mapping.items():
            if not isinstance(source, str) or blank(source) or not isinstance(target, str) or blank(target):
                raise ImportProblem(f"aliases.{name}: название и целевой код должны быть непустыми строками.")
            key = norm(source)
            if key in values and norm(values[key]) != norm(target):
                raise ImportProblem(f"aliases.{name}: для «{source}» заданы разные коды с учётом регистра и пробелов.")
            values[key] = target
        result[name] = values
    return result


def load_config(path=None):
    with open(path or ROOT / "config.json", encoding="utf-8") as stream:
        config = json.load(stream)
    keys = config["group_key"] + config["row_key"]
    if len(keys) != len(set(keys)):
        raise ImportProblem("В настройке ключа повторяется одно из полей.")
    if any(name in keys for name in ("IFT", "PROM", "VALUE", "TEXT_FOR_CLIENT")):
        raise ImportProblem("Результаты решения и IFT/PROM нельзя включать в ключ записи.")
    configured_aliases(config)
    return config


def sheet_named(book, name):
    matches = [s for s in book if norm(s.title) == norm(name)]
    # Excel ограничивает имя листа 31 символом; Р7 может сохранять более длинные имена.
    if not matches and len(name) > 31:
        matches = [s for s in book if norm(s.title) == norm(name[:31])]
    if len(matches) != 1:
        raise ImportProblem(f"Не найден единственный лист «{name}».")
    return matches[0]


def columns_at(sheet, row):
    result = {}
    for cells in sheet.iter_rows(min_row=row, max_row=row):
        for cell in cells:
            if not blank(cell.value):
                key = header(cell.value)
                if key in result:
                    raise ImportProblem(f"{sheet.title}, строка {row}: повторный столбец {cell.value}.")
                result[key] = cell.column
    return result


def literal(cell):
    if cell.data_type == "f":
        raise ImportProblem(f"{cell.parent.title}!{cell.coordinate}: неожиданная формула. "
                            "Для этого поля требуется исходное значение.")
    if cell.data_type == "e":
        raise ImportProblem(f"{cell.parent.title}!{cell.coordinate}: ошибка {cell.value}.")
    return cell.value


@dataclass(frozen=True)
class Entry:
    code: object
    description: object


class Catalog:
    def __init__(self, book, config):
        self.book = book
        self.config = config
        self.indices = {}
        self.invalid_aliases = {}
        self.issues = {}
        self.catalog_locations = {}
        self.aliases = configured_aliases(config)
        self.cache = {}

    def _build_index(self, name):
        spec = self.config["lookups"][name]
        try:
            sheet = sheet_named(self.book, spec["sheet"])
        except ImportProblem as exc:
            raise ImportProblem(f"Основная матрица, справочник «{spec['sheet']}»: {exc}") from exc
        index, invalid = defaultdict(set), defaultdict(list)
        issues = []
        for row in sheet.iter_rows(min_row=2, max_col=max(spec["code_col"], spec["description_col"])):
            code_cell, desc_cell = row[spec["code_col"] - 1], row[spec["description_col"] - 1]
            values, problems = [], []
            for label, cell in (("код", code_cell), ("описание", desc_cell)):
                try:
                    value = template_value(cell)
                except ImportProblem as exc:
                    value = None
                    problems.append(str(exc))
                else:
                    if blank(value):
                        problems.append(f"{cell.coordinate}: пустое поле «{label}»")
                values.append(value)
            code, desc = values
            if blank(code) and blank(desc) and all(blank(c.value) for c in (code_cell, desc_cell)):
                continue
            if problems:
                issue = f"строка {code_cell.row}, {code_cell.coordinate}/{desc_cell.coordinate}: " + "; ".join(problems)
                issues.append(issue)
                # Ошибка блокирует поиск только по связанному с ней коду/описанию.
                # Неразрешённые формулы не превращаем в нули и не используем их кэш.
                for alias in {norm(v) for v in values if not blank(v)}:
                    invalid[alias].append(issue)
                continue
            entry = Entry(code, desc)
            index[norm(code)].add(entry)
            index[norm(desc)].add(entry)
        self.indices[name] = index
        self.invalid_aliases[name] = invalid
        self.issues[name] = issues
        self.catalog_locations[name] = (sheet.title, get_column_letter(spec["code_col"]), get_column_letter(spec["description_col"]))

    def _resolve_direct(self, name, value):
        if name not in self.indices:
            self._build_index(name)
        title, code_col, desc_col = self.catalog_locations[name]
        invalid = self.invalid_aliases[name].get(norm(value), [])
        if invalid:
            details = "; ".join(invalid[:3])
            raise ImportProblem(f"{name}: для значения «{value}» в основной матрице, справочник «{title}», есть неполная или непрочитанная запись: {details}.")
        candidates = self.indices[name].get(norm(value), set())
        if len(candidates) != 1:
            reason = "не найдено" if not candidates else "соответствует нескольким записям"
            details = ""
            if not candidates and self.issues[name]:
                details = " Неполные/непрочитанные строки, которые можно проверить: " + "; ".join(self.issues[name][:3]) + "."
            error = AmbiguousLookup if candidates else ImportProblem
            raise error(f"{name}: значение «{value}» {reason} в основной матрице, справочник «{title}» "
                                f"(код — столбец {code_col}, описание — {desc_col}).{details}")
        return next(iter(candidates))

    def resolve(self, name, value):
        cache_key = (name, norm(value))
        if cache_key in self.cache:
            return self.cache[cache_key]
        target = self.aliases.get(name, {}).get(norm(value))
        if target is None:
            entry = self._resolve_direct(name, value)
        else:
            try:
                entry = self._resolve_direct(name, target)
            except ImportProblem as exc:
                raise ImportProblem(f"aliases.{name}: соответствие «{value}» → «{target}» не удалось применить. {exc}") from exc
            if norm(entry.code) != norm(target):
                raise ImportProblem(f"aliases.{name}: «{target}» — описание, а требуется технический код из основной матрицы.")
            # Существующий однозначный код/описание нельзя незаметно переназначить.
            original = self.indices[name].get(norm(value), set())
            if original and original != {entry}:
                raise ImportProblem(f"aliases.{name}: соответствие «{value}» → «{target}» конфликтует с записью основной матрицы.")
            invalid = self.invalid_aliases[name].get(norm(value), [])
            if invalid:
                raise ImportProblem(f"aliases.{name}: для исходного значения «{value}» есть неполная запись основной матрицы: {invalid[0]}.")
        self.cache[cache_key] = entry
        return entry

    def accepts(self, name, value, entry):
        """Известный код может уточнить неоднозначное описание, но не заменить его."""
        if norm(value) in self.aliases.get(name, {}):
            return self.resolve(name, value) == entry
        if name not in self.indices:
            self._build_index(name)
        if self.invalid_aliases[name].get(norm(value)):
            self._resolve_direct(name, value)  # Точная диагностика неполной записи.
        return entry in self.indices[name].get(norm(value), set())

    def possible_codes(self, name, value):
        """None — неизвестное соответствие; его нельзя исключать при поиске конфликтов."""
        try:
            if norm(value) in self.aliases.get(name, {}):
                return {norm(self.resolve(name, value).code)}
            if name not in self.indices:
                self._build_index(name)
            if self.invalid_aliases[name].get(norm(value)):
                return None
            entries = self.indices[name].get(norm(value), set())
            return {norm(entry.code) for entry in entries} or None
        except ImportProblem:
            return None


def amount(value):
    if blank(value):
        return None
    if isinstance(value, bool):
        raise ImportProblem("MAX_AMOUNT должен быть числом, а не логическим значением.")
    try:
        number = Decimal(str(value).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except InvalidOperation as exc:
        raise ImportProblem(f"MAX_AMOUNT: нечисловое значение «{value}».") from exc
    if not number.is_finite() or number < 0:
        raise ImportProblem("MAX_AMOUNT должен быть конечным неотрицательным числом.")
    return int(number) if number == number.to_integral_value() else float(number)


@dataclass
class Record:
    values: dict
    source: str

    def key(self, config):
        fields = config["group_key"] + config["row_key"]
        parts = [norm(self.values.get(field)) for field in fields]
        platform = self.values.get("SBOL_PRO_PLATFORM") if norm(self.values.get("CHANNEL_INDICATOR")) == norm(config["platform_channel"]) else None
        return tuple(parts + [norm(platform)])


def canonical(raw, catalog, config, source, require_payload=True):
    values = dict(raw)
    required = set(config["group_key"] + ["STOP_LIST_STATUS", "MATCH_NEGATIVE_CASE", "INSPECTED_PERSON"])
    if require_payload:
        required.add("VALUE")
    for name, spec in config["lookups"].items():
        if name == "VALUE" and not require_payload:
            continue
        description = spec["description_field"]
        code_value = raw.get(name)
        desc_value = raw.get(description) if name != "VALUE" else None
        value = code_value if not blank(code_value) else desc_value
        if name == "TRANSACTION_TYPE" and blank(value):
            value = transaction_default(config)
        if blank(value):
            if name in required:
                raise ImportProblem(f"Не заполнено обязательное поле {name}.")
            values[name] = None
            values[description] = None
            continue
        entry = catalog.resolve(name, value)
        if not blank(code_value) and not blank(desc_value):
            if not catalog.accepts(name, desc_value, entry):
                raise ImportProblem(f"{name}: код «{code_value}» не соответствует описанию «{desc_value}».")
        if require_payload and name != "VALUE" and len(catalog.indices[name].get(norm(entry.description), set())) > 1:
            # При записи описания сохранённая Excel-формула могла бы выбрать
            # другой код. Явный входной код не делает такую запись безопасной.
            raise AmbiguousLookup(f"{name}: описание «{entry.description}» соответствует нескольким записям справочника основной матрицы; формульный код нельзя гарантировать.")
        values[name] = entry.code
        values[description] = entry.description
    if norm(values.get("CHANNEL_INDICATOR")) == norm(config["platform_channel"]) and blank(values.get("SBOL_PRO_PLATFORM")):
        raise ImportProblem("Для канала SBOL_PRO требуется SBOL_PRO_PLATFORM.")
    values["MAX_AMOUNT"] = amount(raw.get("MAX_AMOUNT"))
    return Record(values, source)


TOP_CELLS = {
    "CHANNEL_INDICATOR_DESCRIPTION": ("D3", "B3"),
    "SP_NAME_DESCRIPTION": ("D4", "B4"),
    "SP_NAME": ("D5", "B5"),
    "OPERATION_DESCRIPTION": ("D6", "B6"),
    "OPERATION_TYPE": ("D7", "B7"),
    "CLIENT_TYPE_DESCRIPTION": ("B8",),
    "RELATION_TYPE_DESCRIPTION": ("B10",),
    "RELATION_TYPE": ("D10",),
}


CELL_LINK = re.compile(
    r"\s*=\s*(?:(?:'(?P<quoted>(?:[^']|'')+)'|(?P<plain>[^'!\[\]+*/(),;:=<>]+))!)?"
    r"(?P<address>\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)\s*"
)


def template_value(cell, visited=None):
    """Читает значение или простую ссылку, не полагаясь на кэш пересчёта Excel."""
    if cell.data_type != "f":
        return literal(cell)
    visited = set() if visited is None else visited
    location = (cell.parent.title, cell.coordinate)
    if location in visited:
        raise ImportProblem(f"{cell.parent.title}!{cell.coordinate}: циклическая ссылка в шаблоне.")
    if len(visited) >= 64:
        raise ImportProblem(f"{cell.parent.title}!{cell.coordinate}: слишком длинная цепочка ссылок.")
    match = CELL_LINK.fullmatch(cell.value) if isinstance(cell.value, str) else None
    if match is None:
        raise ImportProblem(f"{cell.parent.title}!{cell.coordinate}: неожиданная формула. "
                            "Поддерживаются простые ссылки на ячейки; для вычисления этой формулы требуется отдельное правило.")
    name = match.group("quoted")
    if name is not None:
        name = name.replace("''", "'")
    elif match.group("plain") is not None:
        name = match.group("plain").strip()
    sheet = cell.parent if name is None else sheet_named(cell.parent.parent, name)
    address = match.group("address").replace("$", "").upper()
    try:
        target = sheet[address]
    except (ValueError, IndexError) as exc:
        raise ImportProblem(f"{cell.parent.title}!{cell.coordinate}: недопустимая ссылка.") from exc
    return template_value(target, visited | {location})


def parse_template(path, catalog, config, conflicts=None, outcomes=None):
    book = load_workbook(path, read_only=False, data_only=False)
    try:
        matches = [s for s in book if norm(config["template_sheet_contains"]) in norm(s.title)]
        if len(matches) != 1:
            raise ImportProblem("Не найден единственный лист шаблона с «ЗАПОЛНИТЬ» в имени.")
        sheet = matches[0]
        cols, header_row = None, None
        marker_fields = {"STOP_LIST_STATUS_DESCRIPTION", "MATCH_NEGATIVE_CASE_DESCRIPTION", "VALUE"}
        for row in range(1, min(sheet.max_row, config["template_header_scan_rows"]) + 1):
            row_values = [header(c.value) for c in sheet[row]]
            if marker_fields.issubset(row_values):
                cols, header_row = columns_at(sheet, row), row
                break
        if cols is None:
            raise ImportProblem("Не найдены технические заголовки STOP_LIST_STATUS_DESCRIPTION, MATCH_NEGATIVE_CASE_DESCRIPTION, VALUE.")
        missing = {"MAX_AMOUNT", "TITLE_FOR_CLIENT", "TEXT_FOR_CLIENT"} - cols.keys()
        if missing:
            raise ImportProblem("В шаблоне отсутствуют столбцы: " + ", ".join(sorted(missing)))
        common = {}

        def common_value(field):
            if field not in common:
                common[field] = None
                addresses = TOP_CELLS.get(field, ())
                # Код в формульной D8 определяем по B8 через справочник.
                if field == "CLIENT_TYPE" and sheet["D8"].data_type != "f":
                    addresses = ("D8",)
                for address in addresses:
                    value = template_value(sheet[address])
                    if not blank(value):
                        common[field] = value
                        break
            return common[field]

        def fill_common(raw):
            for name, spec in config["lookups"].items():
                if name == "VALUE":
                    continue
                description = spec["description_field"]
                # Заполненные значения строки приоритетнее всей шапки.
                if not blank(raw.get(name)) or not blank(raw.get(description)):
                    continue
                value = common_value(description)
                if not blank(value):
                    raw[description] = value
                else:
                    value = common_value(name)
                    if not blank(value):
                        raw[name] = value
        result = []
        row_count = 0
        known = {header(k): k for k in config["lookups"]}
        known.update({header(v["description_field"]): v["description_field"] for k, v in config["lookups"].items() if k != "VALUE"})
        for field in ("VALUE", "MAX_AMOUNT", "TITLE_FOR_CLIENT", "TEXT_FOR_CLIENT"):
            known[header(field)] = field
        for row in range(max(header_row + 1, config["template_data_start_row"]), sheet.max_row + 1):
            value_col = cols["VALUE"]
            signals = [sheet.cell(row, cols[x]).value for x in marker_fields]
            if all(blank(x) for x in signals):
                continue
            row_count += 1
            source = f"{Path(path).name}, {sheet.title}!{row}"
            try:
                raw = {field: template_value(sheet.cell(row, cols[key])) for key, field in known.items() if key in cols}
                fill_common(raw)
                raw["VALUE"] = template_value(sheet.cell(row, value_col))
                raw["TITLE_FOR_CLIENT"] = raw.get("TITLE_FOR_CLIENT")
                raw["TEXT_FOR_CLIENT"] = raw.get("TEXT_FOR_CLIENT")
                result.append(canonical(raw, catalog, config, source))
            except AmbiguousLookup as exc:
                if conflicts is None:
                    raise
                reason = f"{exc} Строка пропущена; требуется ручной разбор."
                conflicts.append(f"{source}: {reason}")
                if outcomes is not None:
                    outcomes.append(RecordOutcome(source, "conflict", [], reason))
            except ImportProblem as exc:
                raise ImportProblem(f"{source}: {exc}") from exc
        if not row_count:
            raise ImportProblem("В шаблоне нет строк с решениями.")
        return result
    finally:
        book.close()


def read_matrix_row(sheet, cols, row, catalog, config):
    raw = {}
    for field in config["group_key"] + config["row_key"] + ["SBOL_PRO_PLATFORM"]:
        if field in config["lookups"]:
            desc = config["lookups"][field]["description_field"]
            for f in (field, desc):
                if header(f) in cols:
                    cell = sheet.cell(row, cols[header(f)])
                    if cell.data_type != "f":
                        raw[f] = literal(cell)
                    elif field == "TRANSACTION_TYPE" and f == desc:
                        # Формула — не пустая ячейка. Не подменяем её значением по умолчанию.
                        raw[f] = template_value(cell)
        elif header(field) in cols:
            raw[field] = literal(sheet.cell(row, cols[header(field)]))
    return canonical(raw, catalog, config, f"{sheet.title}!{row}", require_payload=False)


def uncertain_constraints(sheet, cols, row, catalog, config):
    """Возможные ключи неразобранной строки. Ошибка не считается отсутствием записи."""
    constraints = {}
    required = set(config["group_key"] + ["STOP_LIST_STATUS", "MATCH_NEGATIVE_CASE", "INSPECTED_PERSON"])
    for name in dict.fromkeys(config["group_key"] + config["row_key"] + ["SBOL_PRO_PLATFORM"]):
        fields = [name]
        if name in config["lookups"]:
            fields.append(config["lookups"][name]["description_field"])
        values, unreadable, has_formula = [], False, False
        for f in fields:
            if header(f) not in cols:
                continue
            cell = sheet.cell(row, cols[header(f)])
            if cell.data_type == "f":
                has_formula = True
                if name == "TRANSACTION_TYPE" and f != name:
                    try:
                        value = template_value(cell)
                        if not blank(value):
                            values.append(value)
                    except ImportProblem:
                        unreadable = True
                continue
            try:
                value = literal(cell)
                if not blank(value):
                    values.append(value)
            except ImportProblem:
                unreadable = True
        if unreadable:
            choices = None
        elif not values and name == "TRANSACTION_TYPE":
            choices = catalog.possible_codes(name, transaction_default(config))
        elif not values:
            choices = None if name in required or has_formula else {""}
        elif name in config["lookups"]:
            candidates = [catalog.possible_codes(name, value) for value in values]
            # При противоречивых коде и описании учитываем оба варианта.
            choices = None if any(c is None for c in candidates) else set().union(*candidates)
        else:
            try:
                choices = {norm(amount(v) if name == "MAX_AMOUNT" else v) for v in values}
            except ImportProblem:
                choices = None
        constraints[name] = choices
    return constraints


def could_match(record, constraints, config):
    for name, choices in constraints.items():
        if name == "SBOL_PRO_PLATFORM" and norm(record.values.get("CHANNEL_INDICATOR")) != norm(config["platform_channel"]):
            continue
        if choices is not None and norm(record.values.get(name)) not in choices:
            return False
    return True


def payload(record, cols, config):
    result = {}
    for field, spec in config["lookups"].items():
        if field == "VALUE":
            continue
        desc = spec["description_field"]
        if header(desc) in cols:
            result[desc] = record.values.get(desc)
        if header(field) in cols:
            result[field] = record.values.get(field)
    for field in ("VALUE", "MAX_AMOUNT", "TITLE_FOR_CLIENT", "TEXT_FOR_CLIENT"):
        if field in record.values:
            result[field] = record.values[field]
    return result


def text_write(cell, value):
    cell.value = value
    if isinstance(value, str):
        # Введённый пользователем текст, начинающийся на '=', остаётся текстом.
        cell.data_type = "s"


def grow_ranges(sheet, old_end, new_end, start):
    if new_end <= old_end:
        return
    def extend(rng):
        bounds = range_boundaries(str(rng))
        a, b, c, d = bounds
        return f"{get_column_letter(a)}{b}:{get_column_letter(c)}{new_end}" if old_end <= d < new_end and b <= start else str(rng)
    for table in sheet.tables.values():
        table.ref = extend(table.ref)
        if table.autoFilter:
            table.autoFilter.ref = table.ref
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = extend(sheet.auto_filter.ref)
    for validation in sheet.data_validations.dataValidation:
        validation.sqref = MultiCellRange([CellRange(extend(r)) for r in validation.sqref.ranges])
    conditional = []
    for cf in sheet.conditional_formatting:
        conditional.append((str(cf.sqref), list(sheet.conditional_formatting[cf])))
    for old, rules in conditional:
        new = " ".join(extend(r) for r in old.split())
        if new != old:
            del sheet.conditional_formatting[old]
            for rule in rules:
                sheet.conditional_formatting.add(new, rule)


def append_row(sheet, cols, row, proto, values, config, copy_constants=False):
    for col in range(1, sheet.max_column + 1):
        src, dst = sheet.cell(proto, col), sheet.cell(row, col)
        dst._style = copy(src._style)
        dst.protection = copy(src.protection)
        if src.data_type == "f":
            if not isinstance(src.value, str):
                raise ImportProblem(f"{src.coordinate}: массивная формула не поддерживается прототипом.")
            dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
        elif copy_constants:
            dst.value = copy(src.value)
            dst.data_type = src.data_type
        if copy_constants:
            dst.comment = copy(src.comment)
            if src.hyperlink:
                dst.hyperlink = copy(src.hyperlink)
    sheet.row_dimensions[row].height = sheet.row_dimensions[proto].height
    for field, value in ({} if copy_constants else config["new_row_defaults"]).items():
        if header(field) in cols:
            text_write(sheet.cell(row, cols[header(field)]), value)
    # Если CODE хранится числом, а не формулой, используем последовательность строк.
    if "CODE" in cols and sheet.cell(row, cols["CODE"]).data_type != "f":
        sheet.cell(row, cols["CODE"], row - config["matrix_header_row"])
    for field, value in values.items():
        cell = sheet.cell(row, cols[header(field)])
        if cell.data_type != "f" or field in ("IFT", "PROM"):
            text_write(cell, value)
    # Байпас никогда не переносится из шаблона или старой версии.
    bypass = header(config.get("bypass_target", "VALUE_SWITCH"))
    if bypass in cols:
        text_write(sheet.cell(row, cols[bypass]), None)


def flag_value(cell):
    value = literal(cell)
    normalized = value.strip() if isinstance(value, str) else value
    if blank(value) or (not isinstance(value, bool) and normalized in (0, "0")):
        return 0
    if not isinstance(value, bool) and normalized in (1, "1"):
        return 1
    raise ImportProblem(f"{cell.coordinate}: отметка IFT/PROM должна быть 1, 0 или пустой; найдено «{value}».")


def row_mode(sheet, cols, row, environment):
    if not flag_value(sheet.cell(row, cols[environment])):
        return "inactive"
    if environment == "PROM":
        return "promotion"
    return "version" if flag_value(sheet.cell(row, cols["PROM"])) else "inplace"


def payload_changes(sheet, cols, row, values):
    changes, unsafe = {}, []
    for name in MUTABLE_FIELDS:
        cell = sheet.cell(row, cols[header(name)])
        if cell.data_type in ("f", "e"):
            unsafe.append(name)
            continue
        current, value = cell.value, values[name]
        equal = (blank(current) and blank(value)) or (type(current) == type(value) and current == value)
        if isinstance(current, (int, float)) and not isinstance(current, bool) and isinstance(value, (int, float)) and not isinstance(value, bool):
            equal = current == value
        if not equal:
            changes[name] = value
    return changes, unsafe


@dataclass
class Plan:
    row: int
    values: dict
    prototype: int | None = None
    old_version: int | None = None


@dataclass
class RowChange:
    excel_row: int
    action: str
    changed_fields: list[str]
    related_row: int | None
    first_column_name: str
    first_column_value: str | None
    first_column_kind: str


@dataclass
class RecordOutcome:
    source: str
    status: str
    excel_rows: list[int]
    reason: str


def add_outcome(result, source, status, rows, reason):
    outcome = RecordOutcome(source, status, sorted(set(rows)), reason)
    result.outcomes.append(outcome)
    return outcome


def add_conflict(result, source, reason, rows=()):
    result.conflicts.append(f"{source}: {reason}")
    add_outcome(result, source, "conflict", rows, reason)


def row_change(sheet, row, action, fields, related_row, header_row):
    cell = sheet.cell(row, 1)
    kind = "formula" if cell.data_type == "f" else "error" if cell.data_type == "e" else "blank" if blank(cell.value) else "value"
    return RowChange(row, action, list(fields), related_row,
                     str(sheet.cell(header_row, 1).value or "Первый столбец"),
                     str(cell.value) if kind in ("value", "error") else None, kind)


def format_change_report(result):
    lines = [f"Среда: {'ИФТ' if result.environment == 'IFT' else 'ПРОМ'}.",
             "Ниже указаны номера строк Excel в сохранённом результате, а не значения CODE.",
             "Формулы первого столбца здесь не пересчитываются.", ""]
    if not result.changes:
        lines.append("Изменённых строк нет.")
    else:
        lines += ["Затронуты строки Excel: " + ", ".join(map(str, result.affected_rows)), ""]
    labels = {"added": "добавлена запись", "updated": "изменена запись",
              "version_added": "добавлена ИФТ-версия", "ift_deactivated": "снята отметка IFT",
              "prom_activated": "поставлена отметка PROM", "prom_deactivated": "снята отметка PROM"}
    links = {"version_added": "прежняя строка", "ift_deactivated": "новая версия",
             "prom_activated": "прежняя ПРОМ-строка", "prom_deactivated": "включена строка"}
    for change in result.changes:
        line = f"Строка {change.excel_row}: {labels[change.action]}"
        if change.action == "updated" and change.changed_fields:
            line += "; поля: " + ", ".join(change.changed_fields)
        if change.related_row is not None:
            line += f"; {links[change.action]}: {change.related_row}"
        if change.first_column_kind == "value":
            line += f"; {change.first_column_name}={change.first_column_value}"
        elif change.first_column_kind == "formula":
            line += f"; {change.first_column_name}: значение после пересчёта формулы"
        elif change.first_column_kind == "error":
            line += f"; {change.first_column_name}: ошибка {change.first_column_value}"
        lines.append(line)
    if result.outcomes:
        lines += ["", "Результат по строкам шаблонов:"]
        statuses = {"added": "ДОБАВЛЕНА", "versioned": "НОВАЯ ИФТ-ВЕРСИЯ", "updated": "ИЗМЕНЕНА",
                    "promoted": "ВКЛЮЧЕНА В ПРОМ", "unchanged": "БЕЗ ИЗМЕНЕНИЙ",
                    "skipped": "ПРОПУЩЕНА: ПОВТОР", "conflict": "ПРОПУЩЕНА: РУЧНОЙ РАЗБОР"}
        for outcome in result.outcomes:
            rows = ", ".join(map(str, outcome.excel_rows)) or "не определены"
            lines.append(f"{outcome.source} — {statuses[outcome.status]}; строки матрицы: {rows}. {outcome.reason}")
    return "\n".join(lines)


@dataclass
class Result:
    added: int = 0
    updated: int = 0
    skipped: int = 0
    total: int = 0
    output: str = ""
    conflicts: list[str] = field(default_factory=list)
    versioned: int = 0
    environment: str = "IFT"
    promoted: int = 0
    prom_deactivated: int = 0
    changes: list[RowChange] = field(default_factory=list)
    affected_rows: list[int] = field(default_factory=list)
    change_report: str = ""
    outcomes: list[RecordOutcome] = field(default_factory=list)


def import_files(matrix, templates, output, config_path=None, progress: Callable[[str], None] = lambda _: None, environment="IFT"):
    environment = validate_environment(environment)
    config = load_config(config_path)
    matrix, output = Path(matrix).resolve(), Path(output).resolve()
    templates = [Path(p).resolve() for p in templates]
    if not templates:
        raise ImportProblem("Выберите хотя бы один шаблон.")
    inputs = [matrix, *templates]
    if any(p.suffix.lower() != ".xlsx" for p in [*inputs, output]):
        raise ImportProblem("Первая версия поддерживает только .xlsx.")
    if output in inputs or output.exists():
        raise ImportProblem("Выберите новое имя результата: входные и существующие файлы не перезаписываются.")
    for path in inputs:
        if not path.is_file():
            raise ImportProblem(f"Файл не найден: {path.name}")
    if matrix in templates:
        raise ImportProblem("Основная матрица выбрана как шаблон.")
    progress("Читаю основную матрицу…")
    book = load_workbook(matrix, data_only=False, keep_links=True, rich_text=True)
    try:
        sheet = sheet_named(book, config["matrix_sheet"])
        if sheet.protection.sheet:
            raise ImportProblem("Основной лист защищён от изменений.")
        cols = columns_at(sheet, config["matrix_header_row"])
        needed = {"VALUE", "MAX_AMOUNT", "TITLE_FOR_CLIENT", "TEXT_FOR_CLIENT", "IFT", "PROM"}
        for field, spec in config["lookups"].items():
            if field != "VALUE":
                needed.add(spec["description_field"])
                needed.add(field)
        missing = needed - cols.keys()
        if missing:
            raise ImportProblem("В матрице нет столбцов: " + ", ".join(sorted(missing)))
        catalog = Catalog(book, config)
        incoming, conflicts, outcomes = [], [], []
        for path in templates:
            progress(f"Читаю шаблон {path.name}…")
            try:
                incoming.extend(parse_template(path, catalog, config, conflicts, outcomes))
            except ImportProblem as exc:
                if str(exc).startswith(f"{path.name}, "):
                    raise
                raise ImportProblem(f"{path.name}: {exc}") from exc
        result = Result(total=len(incoming) + len(conflicts), output=str(output), conflicts=conflicts, environment=environment, outcomes=outcomes)
        progress("Сопоставляю записи матрицы…")
        index = defaultdict(list)
        prom_index = defaultdict(list)
        uncertain = []
        modes = {}
        last = config["matrix_header_row"]
        for row in range(last + 1, sheet.max_row + 1):
            # Предзаполненные формулы ниже данных не считаем записями.
            active = [sheet.cell(row, cols[header(f)]).value for f in ("SP_NAME_DESCRIPTION", "OPERATION_DESCRIPTION", "STOP_LIST_STATUS_DESCRIPTION")]
            if all(blank(v) for v in active):
                continue
            last = row
            try:
                if environment == "PROM":
                    ift_active = flag_value(sheet.cell(row, cols["IFT"]))
                    prom_active = flag_value(sheet.cell(row, cols["PROM"]))
                    mode = "promotion" if ift_active or prom_active else "inactive"
                else:
                    mode = row_mode(sheet, cols, row, environment)
                if mode == "inactive":
                    continue
                record = read_matrix_row(sheet, cols, row, catalog, config)
            except ImportProblem as exc:
                uncertain.append((row, uncertain_constraints(sheet, cols, row, catalog, config), str(exc)))
            else:
                if environment == "PROM":
                    if ift_active:
                        index[record.key(config)].append(row)
                    if prom_active:
                        prom_index[record.key(config)].append(row)
                else:
                    index[record.key(config)].append(row)
                modes[row] = mode
            if row % 5000 == 0:
                progress(f"Проверено строк матрицы: {row:,}")
        old_end = last
        plans, seen = [], {}
        # Сначала группируем весь пакет: конфликтующие шаблоны не должны менять
        # строку по принципу «кто первый», даже если один из них появился позднее.
        signatures, template_conflicts = {}, set()
        for record in incoming:
            key = record.key(config)
            values = payload(record, cols, config)
            signature = tuple((k, None if blank(values[k]) else values[k]) for k in MUTABLE_FIELDS)
            if key in signatures and signatures[key] != signature:
                template_conflicts.add(key)
            signatures[key] = signature
        uncertain_matches = {}
        for record in incoming:
            key = record.key(config)
            values = payload(record, cols, config)
            if key in template_conflicts:
                add_conflict(result, record.source, "конфликт шаблонов — для одной записи заданы разные значения. Строка пропущена; требуется ручной разбор.", [*index.get(key, []), *prom_index.get(key, [])])
                continue
            matches = index.get(key, [])
            if key not in uncertain_matches:
                uncertain_matches[key] = [(row, reason) for row, constraints, reason in uncertain if could_match(record, constraints, config)]
            unclear = uncertain_matches[key]
            prom_matches = prom_index.get(key, [])
            if (environment == "IFT" and len(matches) > 1) or len(prom_matches) > 1 or unclear:
                reasons = []
                if len(matches) > 1:
                    reasons.append(f"в матрице несколько подходящих строк{' ИФТ' if environment == 'PROM' else ''}: {matches[:10]}")
                if len(prom_matches) > 1:
                    reasons.append(f"в матрице несколько активных строк ПРОМ: {prom_matches[:10]}")
                if unclear:
                    reasons.append("возможное совпадение с неразобранными строками матрицы: " + "; ".join(f"{row} ({reason})" for row, reason in unclear[:3]))
                add_conflict(result, record.source, f"{'; '.join(reasons)}. Строка пропущена; требуется ручной разбор.", [*matches, *prom_matches, *(r for r, _ in unclear)])
                continue
            if key in seen:
                result.skipped += 1
                first = seen[key]
                add_outcome(result, record.source, "skipped", first.excel_rows, f"Повтор такой же записи в загруженных шаблонах. Первый источник: {first.source}.")
                continue
            if environment == "PROM":
                current_prom = prom_matches[0] if prom_matches else None
                ready, unreadable, mismatch_fields = [], [], set()
                for candidate in matches:
                    candidate_changes, candidate_unsafe = payload_changes(sheet, cols, candidate, values)
                    mismatch_fields.update(candidate_changes)
                    if candidate_unsafe:
                        unreadable.append(candidate)
                    elif not candidate_changes:
                        ready.append(candidate)
                if len(ready) > 1 or unreadable:
                    reason = f"несколько ИФТ-строк с одинаковыми ключом и содержимым: {ready[:10]}" if len(ready) > 1 else f"содержимое ИФТ-строк нельзя проверить: {unreadable[:10]}"
                    add_conflict(result, record.source, f"{reason}. Отметки не изменены; требуется ручной разбор.", [*ready, *unreadable, *prom_matches])
                    continue
                already_deployed = False
                if current_prom is not None:
                    current_changes, unsafe = payload_changes(sheet, cols, current_prom, values)
                    already_deployed = not current_changes and not unsafe
                    if already_deployed and (not ready or ready[0] == current_prom):
                        seen[key] = add_outcome(result, record.source, "unchanged", [current_prom], "Нужные реакция, заголовок и текст уже действуют в ПРОМ; переключение отметок не требуется.")
                        result.skipped += 1
                        continue
                if not ready:
                    reason = "нет активной ИФТ-строки" if not matches else f"содержимое активных ИФТ-строк {matches[:10]} не совпадает с шаблоном"
                    if mismatch_fields:
                        reason += "; отличаются поля: " + ", ".join(sorted(mismatch_fields))
                    add_conflict(result, record.source, f"{reason} для включения в ПРОМ. Ничего не изменено; требуется ручной разбор.", [*matches, *prom_matches])
                    continue
                row = ready[0]
                # В ПРОМ переносится готовая ИФТ-версия. Её данные и байпас
                # не переписываются, новые строки здесь не создаются.
                plans.append(Plan(row, {"PROM": 1}, old_version=current_prom))
                rows = [row] + ([current_prom] if current_prom is not None else [])
                seen[key] = add_outcome(result, record.source, "promoted", rows, f"Готовой ИФТ-строке {row} поставлен PROM=1." + (f" У прежней строки {current_prom} снят только PROM." if current_prom is not None else " Прежней ПРОМ-версии нет."))
                result.updated += 1
                result.promoted += 1
                if current_prom is not None:
                    result.prom_deactivated += 1
                continue
            if matches:
                row = matches[0]
                # Ключи и прочие столбцы существующей записи не редактируются.
                changes, unsafe = payload_changes(sheet, cols, row, values)
                if unsafe:
                    add_conflict(result, record.source, f"поля {', '.join(unsafe)} в строке матрицы {row} содержат формулу или ошибку. Строка пропущена; требуется ручной разбор.", [row])
                    continue
                if changes:
                    if modes[row] == "version":
                        last += 1
                        version_values = dict(changes)
                        for field in ("TRANSACTION_TYPE", config["lookups"]["TRANSACTION_TYPE"]["description_field"]):
                            cell = sheet.cell(row, cols[header(field)])
                            if cell.data_type != "f" and blank(cell.value):
                                version_values[field] = record.values.get(field)
                        version_values.update(IFT=1 if environment == "IFT" else None, PROM=1 if environment == "PROM" else None)
                        plans.append(Plan(last, version_values, prototype=row, old_version=row))
                        result.added += 1
                        result.versioned += 1
                        seen[key] = add_outcome(result, record.source, "versioned", [row, last], f"Добавлена ИФТ-версия {last}; у прежней строки {row} снят только IFT. Изменены поля: {', '.join(changes)}.")
                    else:
                        plans.append(Plan(row, changes))
                        result.updated += 1
                        seen[key] = add_outcome(result, record.source, "updated", [row], "Изменены поля: " + ", ".join(changes) + ".")
                else:
                    result.skipped += 1
                    seen[key] = add_outcome(result, record.source, "unchanged", [row], "Реакция, экранный заголовок и текст совпадают; изменений не требуется.")
            else:
                last += 1
                values.update(IFT=1 if environment == "IFT" else None, PROM=1 if environment == "PROM" else None)
                plans.append(Plan(last, values, prototype=config["prototype_row"]))
                result.added += 1
                seen[key] = add_outcome(result, record.source, "added", [last], "Активная ИФТ-запись с таким ключом отсутствовала; добавлена новая строка.")
        proto = config["prototype_row"]
        if proto <= config["matrix_header_row"] or proto > old_end:
            raise ImportProblem("В матрице нет образцовой строки для копирования формул и оформления.")
        progress("Записываю изменения и сохраняю новый файл…")
        for plan in plans:
            row, values = plan.row, plan.values
            if plan.prototype is not None:
                if any(not blank(c.value) and c.data_type != "f" for c in sheet[row]):
                    raise ImportProblem(f"Строка {row} ниже таблицы содержит данные. Добавление поверх неё остановлено.")
                for field_name in values:
                    if field_name not in config["lookups"] and field_name not in ("IFT", "PROM") and sheet.cell(plan.prototype, cols[header(field_name)]).data_type == "f":
                        raise ImportProblem(f"Образцовая строка {plan.prototype}: входное поле {field_name} содержит формулу.")
                append_row(sheet, cols, row, plan.prototype, values, config, copy_constants=plan.old_version is not None)
                action = "version_added" if plan.old_version is not None else "added"
                result.changes.append(row_change(sheet, row, action, [], plan.old_version, config["matrix_header_row"]))
            else:
                for field, value in values.items():
                    text_write(sheet.cell(row, cols[header(field)]), value)
                action = "prom_activated" if environment == "PROM" else "updated"
                result.changes.append(row_change(sheet, row, action, values.keys(), plan.old_version, config["matrix_header_row"]))
            if plan.old_version is not None:
                text_write(sheet.cell(plan.old_version, cols[environment]), None)
                action = "prom_deactivated" if environment == "PROM" else "ift_deactivated"
                result.changes.append(row_change(sheet, plan.old_version, action, [environment], row, config["matrix_header_row"]))
        grow_ranges(sheet, old_end, last, config["matrix_header_row"] + 1)
        book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto")
        output.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(prefix=".matrix-import-", suffix=".xlsx", dir=output.parent)
        os.close(fd)
        try:
            book.save(tmp)
            # Эксклюзивное создание не позволяет затереть файл, появившийся за время работы.
            with output.open("xb") as dest, open(tmp, "rb") as src:
                import shutil
                try:
                    shutil.copyfileobj(src, dest)
                except BaseException:
                    dest.close()
                    output.unlink(missing_ok=True)
                    raise
        finally:
            Path(tmp).unlink(missing_ok=True)
        result.changes.sort(key=lambda change: change.excel_row)
        result.affected_rows = sorted({change.excel_row for change in result.changes})
        result.change_report = format_change_report(result)
        return result
    finally:
        book.close()
